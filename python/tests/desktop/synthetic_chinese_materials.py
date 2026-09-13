"""Chinese acceptance variant only; the shared eight-file and generic five-file oracles stay unchanged.

All people and clauses are fictional source statements, not expected model findings.
Run this module with a uniquely owned empty /private/tmp directory to generate inputs.
"""
from datetime import datetime
from io import BytesIO
import hashlib
import json
from pathlib import Path
import sys

from docx import Document
from openpyxl import Workbook
import pymupdf

from synthetic_mixed_materials import EMPLOYEES, saved

OCR_LINES = (
    '完全虚构验收材料 — 解除通知及签收记录',
    '企业：合成星河制造有限公司',
    '员工：SYN-010 合成员工十',
    'SYN-010 解除日期：2026-09-01',
    '双方协商结束劳动关系，办理离职交接。',
    '送达记录：2026-09-01 本人签收。',
    '本页仅是合成材料，不含真实个人资料。',
)
NAMES = ('一', '二', '三', '四', '五', '六', '七', '八', '九', '十')


def contract_line(index, eid):
    return f'{eid} 合成员工{NAMES[index]}：双方于2026-01-01签署固定期限劳动合同。'


def wage_line(index, eid):
    clause = '单位可不经通知扣留全部工资。' if index % 2 == 0 else '工资每月十五日支付，变更须双方书面确认。'
    return f'{eid} 工资条款：{clause}'


def build_materials():
    materials = {'01-员工花名册.csv': ('员工编号,姓名,入职日期,部门\n' + '\n'.join(
        f'{eid},合成员工{NAMES[i]},2026-01-01,合成制造部' for i, eid in enumerate(EMPLOYEES))).encode('utf-8-sig')}
    doc = Document(); doc.add_paragraph('合成星河制造有限公司：劳动合同合成验收材料')
    for i, eid in enumerate(EMPLOYEES): doc.add_paragraph(contract_line(i, eid))
    table = doc.add_table(rows=0, cols=2)
    for eid in EMPLOYEES:
        cells = table.add_row().cells; cells[0].text = eid
        cells[1].text = '合同期限：2026-01-01至2027-01-01'
    for i, eid in enumerate(EMPLOYEES): doc.add_paragraph(wage_line(i, eid))
    materials['02-劳动合同与工资条款.docx'] = saved(doc)
    workbook = Workbook(); sheet = workbook.active; sheet.title = '试用期记录'
    sheet.append(['员工编号', '试用期开始', '试用期结束', '考核记录'])
    for i, eid in enumerate(EMPLOYEES):
        sheet.append([eid, datetime(2026, 1, 1), datetime(2026, 3, 1), '已留存考核表' if i % 2 else '尚未提供考核表'])
    materials['03-试用期记录.xlsx'] = saved(workbook)
    with pymupdf.open() as pdf:
        for group in (EMPLOYEES[:5], EMPLOYEES[5:]):
            page = pdf.new_page(width=700, height=500)
            for i, eid in enumerate(group):
                page.insert_text((35, 60 + i * 60), f'{eid} 续签通知：拟于2027-01-01续订劳动合同。', fontname='china-s', fontsize=20)
        materials['04-续签通知.pdf'] = pdf.tobytes()
    materials['05-社保记录.csv'] = ('员工编号,社保月份,缴纳主体,原文备注\n' + '\n'.join(
        f'{eid},2026-08,合成星河制造有限公司,' + ('本人自愿放弃社会保险。' if i == 2 else '已提供本月缴费记录。')
        for i, eid in enumerate(EMPLOYEES))).encode('utf-8-sig')
    with pymupdf.open() as page_document:
        page = page_document.new_page(width=700, height=460)
        for i, line in enumerate(OCR_LINES):
            page.insert_text((35, 58 + i * 52), line, fontname='china-s', fontsize=22)
        png = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False).tobytes('png')
    materials['06-解除通知.png'] = png
    with pymupdf.open() as scan:
        page = scan.new_page(width=700, height=460); page.insert_image(page.rect, stream=png)
        materials['07-解除通知扫描.pdf'] = scan.tobytes()
    doc = Document(); doc.add_paragraph('SYN-001 合成合同签署页：2026-01-01，另附嵌入图片，图片内容待单独核对。')
    doc.add_picture(BytesIO(png)); materials['08-嵌图合同待核对.docx'] = saved(doc)
    return materials


def observations():
    rows = []
    def add(filename, ids, location, text):
        rows.append({'filename': filename, 'employee_ids': ids, 'location': location, 'text': text})
    for i, eid in enumerate(EMPLOYEES):
        add('01-员工花名册.csv', [eid], {'sheet': 'CSV', 'row': i + 2}, f'{eid} | 合成员工{NAMES[i]} | 2026-01-01')
        add('02-劳动合同与工资条款.docx', [eid], {'paragraph': i + 2}, contract_line(i, eid))
        add('02-劳动合同与工资条款.docx', [eid], {'table': 1, 'row': i + 1, 'column': 2}, '合同期限：2026-01-01至2027-01-01')
        add('02-劳动合同与工资条款.docx', [eid], {'paragraph': i + 12}, wage_line(i, eid))
        add('03-试用期记录.xlsx', [eid], {'sheet': '试用期记录', 'row': i + 2, 'column': 3}, '2026-03-01')
        add('04-续签通知.pdf', [eid], {'page': 1 if i < 5 else 2}, f'{eid} 续签通知：拟于2027-01-01续订劳动合同。')
        add('05-社保记录.csv', [eid], {'sheet': 'CSV', 'row': i + 2}, f'{eid} | 2026-08 | 合成星河制造有限公司')
    for filename in ('06-解除通知.png', '07-解除通知扫描.pdf'):
        for i, text in enumerate(OCR_LINES): add(filename, ['SYN-010'], {'page': 1, 'line': i + 1}, text)
    add('08-嵌图合同待核对.docx', ['SYN-001'], {'paragraph': 1}, 'SYN-001 合成合同签署页：2026-01-01')
    return rows


def generate(directory):
    target = Path(directory).resolve()
    if target.parent != Path('/private/tmp') or not target.is_dir() or any(target.iterdir()):
        raise ValueError('Use a uniquely owned empty /private/tmp directory')
    materials = build_materials()
    for name, content in materials.items(): (target / name).write_bytes(content)
    (target / 'source-observations.json').write_text(json.dumps(observations(), ensure_ascii=False, indent=2))
    manifest = [{'filename': name, 'size': len(data), 'sha256': hashlib.sha256(data).hexdigest()} for name, data in materials.items()]
    (target / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2))
    return manifest


if __name__ == '__main__':
    print(json.dumps(generate(sys.argv[1]), ensure_ascii=False, indent=2))
