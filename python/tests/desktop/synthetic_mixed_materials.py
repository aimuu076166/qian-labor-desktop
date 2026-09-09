"""Reusable ten-employee inputs, with no expected findings or personal identifiers.

build_materials() returns filename -> bytes for offline or separately authorized
real-provider/installed-app acceptance. All people and statements are fictional.
"""
from datetime import datetime
from io import BytesIO

import pymupdf
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw

EMPLOYEES = tuple(f"SYN-{number:03}" for number in range(1, 11))
OCR_LINE = "SYN-010 termination notice 2026-09-01"


def saved(document):
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def build_materials():
    materials = {"roster.csv": ("employee,name,start\n" + "\n".join(
        f"{eid},Fictional employee {i},2026-01-01" for i, eid in enumerate(EMPLOYEES, 1))).encode()}
    doc = Document()
    for eid in EMPLOYEES:
        doc.add_paragraph(f"{eid} signed contract 2026-01-01")
    table = doc.add_table(rows=0, cols=2)
    for eid in EMPLOYEES:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = eid, "Contract term 2026-01-01 to 2027-01-01"
    for index, eid in enumerate(EMPLOYEES):
        clause = "Employer may withhold all wages without notice." if index % 2 == 0 else "Wages paid monthly on the date agreed in writing."
        doc.add_paragraph(f"{eid} Wage clause: {clause}")
    materials["contracts.docx"] = saved(doc)
    workbook = Workbook()
    workbook.active.title = "Probation"
    workbook.active.append(["employee", "probation_end"])
    for eid in EMPLOYEES:
        workbook.active.append([eid, datetime(2026, 3, 1)])
    materials["probation.xlsx"] = saved(workbook)
    pdf = pymupdf.open()
    for group in (EMPLOYEES[:5], EMPLOYEES[5:]):
        page = pdf.new_page()
        for i, eid in enumerate(group):
            page.insert_text((40, 50 + i * 25), f"{eid} renewal notice 2027-01-01")
    materials["renewal.pdf"] = pdf.tobytes()
    pdf.close()
    materials["social.csv"] = ("employee,social_record\n" + "\n".join(f"{eid},2026-08 contribution record" for eid in EMPLOYEES)).encode()
    image = Image.new("RGB", (600, 80), "white")
    ImageDraw.Draw(image).text((10, 20), OCR_LINE, fill="black")
    out = BytesIO()
    image.save(out, format="PNG")
    materials["termination.png"] = out.getvalue()
    scan = pymupdf.open()
    page = scan.new_page(width=600, height=80)
    page.insert_image(page.rect, stream=out.getvalue())
    materials["termination-scan.pdf"] = scan.tobytes()
    scan.close()
    doc = Document()
    doc.add_paragraph("SYN-001 signed contract")
    doc.add_picture(BytesIO(out.getvalue()))
    materials["embedded-warning.docx"] = saved(doc)
    return materials
