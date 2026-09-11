"""Offline synthetic contract observations; no expected production legal conclusions."""
import json
from io import BytesIO

import pytest
from docx import Document
from pydantic import ValidationError

from qian_labor.ai.schemas import ExtractionResult
from qian_labor.ai.grounding import EXTRACTION_VERSION
from qian_labor.jobs.processing import ProcessingPipeline
from qian_labor.parsers.registry import ParserRegistry


def advisory(quote="SYN-001 工资条款：发放日期由双方书面确认。", **position):
    return {"version": "contract-advisory-v1", "status": "completed", "observations": [{
        "employee_number": "SYN-001", "source": {"file_name": "untrusted.docx", "excerpt": quote, **position},
        "issue": "需核对工资发放条款的具体约定", "checks": ["核对签署版本及双方约定"],
        "next_action": "请双方确认发放日期并保留书面记录。", "unverified_references": [],
    }]}


def contract_input():
    doc = Document()
    doc.add_paragraph("SYN-001 工资条款：发放日期由双方书面确认。")
    doc.add_paragraph("SYN-001 工资条款：发放日期由双方书面确认。")
    out = BytesIO()
    doc.save(out)
    content = out.getvalue()
    return ProcessingPipeline._extraction_inputs("contract.docx", content, ParserRegistry().parse("contract.docx", content))[0]


def test_optional_typed_advisory_preserves_legacy_and_clause_only_result():
    assert getattr(ExtractionResult(), "contract_advisory", "missing") is None
    result = ExtractionResult.model_validate({"document_type": "contract", "contract_advisory": advisory()})
    assert result.facts == []
    assert result.contract_advisory.version == "contract-advisory-v1"


@pytest.mark.parametrize("field,value", [("issue", "x" * 1001), ("checks", ["x"] * 9), ("next_action", ""), ("unverified_references", ["x"] * 6)])
def test_advisory_rejects_invalid_bounds_without_truncating(field, value):
    assert ExtractionResult.model_validate({"contract_advisory": advisory()}).contract_advisory is not None
    payload = advisory()
    payload["observations"][0][field] = value
    with pytest.raises(ValidationError):
        ExtractionResult.model_validate({"contract_advisory": payload})


def test_advisory_grounding_is_separate_preserves_multiple_real_positions():
    from qian_labor.ai import grounding
    assert hasattr(grounding, "ground_advisory"), "Need typed local advisory grounding, never placeholder employment facts"
    result = ExtractionResult.model_validate({"contract_advisory": advisory()})
    rows = grounding.ground_advisory(result.contract_advisory, contract_input(), "contract.docx")
    assert result.facts == []
    assert [row.source.paragraph for row in rows] == [1, 2]
    assert all(row.proof["requires_review"] for row in rows)
    assert all(row.proof["version"] == EXTRACTION_VERSION for row in rows)
    assert all(row.source.file_name == "contract.docx" for row in rows)


@pytest.mark.parametrize("quote,position", [("虚构的原文", {}), (None, {"paragraph": 99})])
def test_advisory_invented_source_is_unlocated(quote, position):
    from qian_labor.ai import grounding
    assert hasattr(grounding, "ground_advisory")
    result = ExtractionResult.model_validate({"contract_advisory": advisory(**position) if quote is None else advisory(quote)})
    rows = grounding.ground_advisory(result.contract_advisory, contract_input(), "contract.docx")
    assert len(rows) == 1
    assert rows[0].source.excerpt == "" and rows[0].source.paragraph is None
    assert rows[0].proof["status"] == "unlocated_needs_review"


def test_combined_cache_identity_keeps_grounding_version():
    assert ProcessingPipeline._job_key("a", "f", "extract", "sha") == f"a:f:extract:sha:{EXTRACTION_VERSION}:contract-advisory-v1"
    assert ProcessingPipeline._job_key("a", "f", "parse", "sha") == "a:f:parse:sha"


def test_zhipu_clause_only_combines_in_one_configured_request():
    import httpx
    from qian_labor.ai.zhipu_provider import ZhipuChatCompletionsProvider
    from qian_labor.security.local_redaction import PrivacyBoundary, PreparedProviderContent
    requests = []
    payload = {"schema_version": "employment-extraction-v1", "document_type": "contract",
               "employee_name": None, "employee_number": "SYN-001", "department": None,
               "job_title": None, "needs_human_confirmation": True, "facts": [], "contract_advisory": advisory()}
    def respond(request):
        requests.append(request)
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(payload)}}]})
    provider = ZhipuChatCompletionsProvider("synthetic", "https://configured.invalid/custom", "configured-model", "configured-vision",
        privacy_boundary=PrivacyBoundary("synthetic-advisory-pepper-at-least-32-characters"),
        max_attempts=1, client=httpx.Client(transport=httpx.MockTransport(respond)))
    result = provider.extract("contract.txt", PreparedProviderContent(b"SYN-001 synthetic contract"))
    assert result.facts == [] and result.contract_advisory.status == "completed"
    assert len(requests) == 1 and str(requests[0].url) == "https://configured.invalid/custom/chat/completions"
    sent = json.loads(requests[0].content)
    assert sent["model"] == "configured-model"
    assert "unverified" in sent["messages"][0]["content"]


from test_company_workspaces import api, company, record
from test_current_company import current, select_record


class SyntheticClauseProvider:
    name = "fake"
    calls = 0
    def extract(self, filename, content):
        self.calls += 1
        return ExtractionResult(document_type="contract", employee_number="SYN-001", contract_advisory=advisory())


def processed_contract(api, tmp_path):
    from qian_labor.desktop.import_service import DesktopImportService
    client, db = api
    c = company(client)
    rec = record(client, c["id"])
    aid = current(client, c["id"], 1)["analysis_id"]
    doc = Document()
    doc.add_paragraph(advisory()["observations"][0]["source"]["excerpt"])
    path = tmp_path / "contract.docx"
    doc.save(path)
    imports = DesktopImportService(db, tmp_path)
    imports.import_paths(aid, [path])
    provider = SyntheticClauseProvider()
    pipeline = ProcessingPipeline(db, imports.storage, provider=provider)
    pipeline.process(aid)
    return client, db, c, rec, aid, provider, pipeline


def test_clause_only_persists_matching_without_fake_facts_and_handling_cas(api, tmp_path):
    from uuid import uuid4
    from sqlalchemy import select
    from qian_labor.models.core import EmploymentFact
    client, db, c, rec, aid, provider, pipeline = processed_contract(api, tmp_path)
    url = f'/api/company-workspaces/{c["id"]}/analyses/{aid}/contract-advisories'
    response = client.get(url)
    assert response.status_code == 200, response.text
    state = response.json()
    assert state["runs"][0]["execution_status"] == "completed"
    row = state["observations"][0]
    assert row["employee_id"] is None and row["assignment_status"] == "pending_matching"
    assert row["source"]["location"] == {"paragraph": 1}
    with db.session() as s:
        assert list(s.scalars(select(EmploymentFact))) == []
    candidates = client.get(f"/api/analyses/{aid}/matching-candidates").json()["candidates"]
    assert len(candidates) == 1
    assert select_record(client, aid, candidates[0], rec).status_code == 200
    state = client.get(url, params={"record_id": rec["id"]}).json()
    row = state["observations"][0]
    assert row["employee_id"] is not None
    request_id = str(uuid4())
    body = {"id": request_id, "expected_version": 0, "decision": "checking", "reason": "正在核对合成签署版本"}
    response = client.post(f'{url}/{row["id"]}/handling', json=body)
    assert response.status_code == 200, response.text
    assert response.json()["version"] == 1
    assert client.post(f'{url}/{row["id"]}/handling', json=body).status_code == 409
    reconciled = client.get(url, params={"request_id": request_id}).json()
    assert reconciled["request_handling"]["id"] == request_id
    assert reconciled["observations"][0]["source"] == row["source"]
    assert provider.calls == 1
    pipeline.process(aid)
    assert provider.calls == 1


def test_clause_only_workspace_is_not_a_no_supported_facts_error(api, tmp_path):
    client, db, c, rec, aid, provider, pipeline = processed_contract(api, tmp_path)
    file = client.get(f'/api/analyses/{aid}/workspace').json()["files"][0]
    assert file["error_code"] is None
    assert file["advisory_status"] == "completed"
    assert file["fact_count"] == 0


@pytest.mark.parametrize("statuses,expected", [(["completed", "unreadable"], "partial"),
    (["completed", "not_executed"], "partial"), (["not_executed"], "not_executed"),
    (["unreadable"], "unreadable"), (["not_applicable"], "not_applicable"), (["completed"], "completed")])
def test_file_completion_requires_all_inputs(statuses, expected):
    from qian_labor.services.contract_advisory import aggregate_status
    assert aggregate_status(statuses) == expected
    if statuses == ["completed"]:
        assert aggregate_status(statuses, True) == "partial"


def test_cross_company_corrupted_sources_and_historical_handling_are_rejected(api, tmp_path):
    from uuid import uuid4
    from sqlalchemy import select
    from qian_labor.models.core import ContractClauseObservation
    client, db, c, rec, aid, provider, pipeline = processed_contract(api, tmp_path)
    other = company(client)
    url = f'/api/company-workspaces/{c["id"]}/analyses/{aid}/contract-advisories'
    assert client.get(url.replace(c["id"], other["id"])).status_code == 404
    assert client.get(url, params={"record_id": record(client, other["id"])["id"]}).status_code == 404
    row = client.get(url).json()["observations"][0]
    with db.session() as s:
        saved = s.get(ContractClauseObservation, row["id"])
        saved.source_excerpt = "虚构篡改原文"
        s.commit()
    assert client.get(url).json()["detail"]["code"] == "ADVISORY_SOURCE_INVALID"
    body = {"id": str(uuid4()), "expected_version": 0, "decision": "checking", "reason": "合成核对"}
    assert client.post(f'{url}/{row["id"]}/handling', json=body).json()["detail"]["code"] == "ADVISORY_SOURCE_INVALID"
    assert provider.calls == 1


def test_split_sensitive_ocr_tokens_never_persist_in_advisory_excerpt(api, tmp_path):
    from PIL import Image
    from qian_labor.security.local_redaction import LocalImageRedactor, PrivacyBoundary, OCRToken
    from qian_labor.desktop.import_service import DesktopImportService
    # Deliberately invalid 32-digit synthetic opaque identifier, split across OCR tokens.
    parts = ["9" * 8] * 4
    class OCR:
        def extract_tokens(self, content):
            return [OCRToken("SYN-001 contract", 2, 4, 100, 20, "line"),
                    *[OCRToken(part, 110 + i * 40, 4, 35, 20, "line") for i, part in enumerate(parts)]]
    boundary = PrivacyBoundary("synthetic-privacy-pepper-at-least-32-characters", LocalImageRedactor(OCR()))
    out = BytesIO()
    Image.new("RGB", (400, 60), "white").save(out, format="PNG")
    prepared = boundary.prepare("synthetic.png", out.getvalue(), is_image=True, external=True)
    assert all(part not in str(prepared.ocr_blocks) for part in parts)
    client, db = api
    c = company(client)
    aid = current(client, c["id"])["analysis_id"]
    path = tmp_path / "synthetic.png"
    path.write_bytes(out.getvalue())
    imports = DesktopImportService(db, tmp_path)
    imports.import_paths(aid, [path])
    class Provider(SyntheticClauseProvider):
        is_external = True
        def extract(self, filename, content):
            return ExtractionResult(document_type="contract", employee_number="SYN-001",
                contract_advisory=advisory("SYN-001 contract"))
    ProcessingPipeline(db, imports.storage, provider=Provider(), privacy_boundary=boundary).process(aid)
    state = client.get(f'/api/company-workspaces/{c["id"]}/analyses/{aid}/contract-advisories').json()
    excerpt = state["observations"][0]["source"]["excerpt"]
    assert "[REDACTED]" in excerpt and all(part not in excerpt for part in parts)
    from sqlalchemy import select
    from qian_labor.models.core import ContractClauseObservation
    with db.session() as s:
        assert all(part not in row.source_excerpt for row in s.scalars(select(ContractClauseObservation)) for part in parts)


def test_ten_employee_clause_scopes_are_all_reachable_and_only_explicitly_assigned(api, tmp_path):
    from synthetic_mixed_materials import build_materials, EMPLOYEES
    from qian_labor.desktop.import_service import DesktopImportService
    from sqlalchemy import select
    from qian_labor.models.core import EmploymentFact
    client, db = api
    c = company(client)
    records = [record(client, c["id"], number, i) for i, number in enumerate(EMPLOYEES)]
    aid = current(client, c["id"], 10)["analysis_id"]
    data = build_materials()["contracts.docx"]
    blocks = ParserRegistry().parse("contracts.docx", data).blocks
    clauses = [b for b in blocks if "Wage clause:" in b.text]
    assert len(clauses) == 10, "Shared mixed generator needs actual benign and concerning contract clauses"
    class MultiProvider(SyntheticClauseProvider):
        def extract(self, filename, content):
            payload = advisory()
            payload["observations"] = []
            for block in clauses:
                item = advisory(block.text)["observations"][0]
                item["employee_number"] = block.text.split()[0]
                payload["observations"].append(item)
            return ExtractionResult(document_type="contract", contract_advisory=payload)
    path = tmp_path / "contracts.docx"
    path.write_bytes(data)
    imports = DesktopImportService(db, tmp_path)
    imports.import_paths(aid, [path])
    ProcessingPipeline(db, imports.storage, provider=MultiProvider()).process(aid)
    url = f'/api/company-workspaces/{c["id"]}/analyses/{aid}/contract-advisories'
    response = client.get(url, params={"page_size": 3}).json()
    assert response["total"] == 10 and len(response["observations"]) == 3 and response["pages"] == 4
    candidates = client.get(f'/api/analyses/{aid}/matching-candidates').json()["candidates"]
    assert len(candidates) == 10
    for candidate in candidates:
        rec = next(r for r in records if r["employee_number"] == candidate["extracted_fields"]["employee_ids"][0])
        selected = select_record(client, aid, candidate, rec)
        assert selected.status_code == 200, selected.text
        rows = client.get(url, params={"record_id": rec["id"]}).json()["observations"]
        assert len(rows) == 1
        assert rows[0]["source"]["excerpt"].startswith(rec["employee_number"] + " ")
    with db.session() as s:
        assert list(s.scalars(select(EmploymentFact))) == []


def test_explicit_missing_advisory_upgrade_preserves_original_sources_and_old_run(api, tmp_path):
    from sqlalchemy import select
    from qian_labor.models.core import ContractAdvisoryRun, AnalysisBatch, UploadedFile
    from qian_labor.services.contract_advisory import ContractAdvisoryService
    client, db, c, rec, aid, provider, pipeline = processed_contract(api, tmp_path)
    with db.session() as s:
        old = s.scalar(select(ContractAdvisoryRun))
        old.execution_status = "not_executed"
        old.input_statuses = ["not_executed"]
        old_id = old.id
        s.get(AnalysisBatch, aid).status = "completed"
        s.commit()
    state = ContractAdvisoryService(db).listing(c["id"], aid)
    before_row = state["observations"][0]
    pipeline.process(aid)
    assert provider.calls == 1
    pipeline.provider.supports_contract_advisory = True
    client.app.state.processing_queue.pipeline_factory = lambda: pipeline
    assert client.post(f'/api/analyses/{aid}/process').status_code == 202
    client.app.state.processing_queue.shutdown()
    assert provider.calls == 2, "Explicit capability upgrade must not reuse not_executed cache"
    with db.session() as s:
        assert s.get(ContractAdvisoryRun, old_id).execution_status == "not_executed"
        assert len(list(s.scalars(select(ContractAdvisoryRun)))) == 2
    historical = ContractAdvisoryService(db).listing(c["id"], aid, history=True)
    prior = next(r for r in historical["observations"] if r["id"] == before_row["id"])
    assert prior["source"] == before_row["source"] and prior["read_only"]


def test_explicit_unreadable_model_retry_recovers_but_parser_only_partial_dedupes(api, tmp_path):
    from sqlalchemy import select
    from qian_labor.models.core import ContractAdvisoryRun, AnalysisBatch
    client, db, c, rec, aid, provider, pipeline = processed_contract(api, tmp_path)
    pipeline.provider.supports_contract_advisory = True
    with db.session() as s:
        run = s.scalar(select(ContractAdvisoryRun))
        run.execution_status = "partial"
        run.input_statuses = ["completed", "unreadable"]
        s.get(AnalysisBatch, aid).status = "queued"
        s.commit()
    pipeline.process(aid)
    assert provider.calls == 2
    with db.session() as s:
        newest = s.scalar(select(ContractAdvisoryRun).order_by(ContractAdvisoryRun.created_at.desc()))
        newest.execution_status = "partial"
        assert newest.input_statuses == ["completed"]
        s.get(AnalysisBatch, aid).status = "queued"
        s.commit()
    pipeline.process(aid)
    assert provider.calls == 2


def test_handling_restart_masking_and_scoped_request_lookup(api, tmp_path):
    from uuid import uuid4
    from qian_labor.desktop.app import create_desktop_app
    from fastapi.testclient import TestClient
    from test_company_workspaces import HEADERS
    client, db, c, rec, aid, provider, pipeline = processed_contract(api, tmp_path)
    candidate = client.get(f'/api/analyses/{aid}/matching-candidates').json()["candidates"][0]
    assert select_record(client, aid, candidate, rec).status_code == 200
    url = f'/api/company-workspaces/{c["id"]}/analyses/{aid}/contract-advisories'
    row = client.get(url).json()["observations"][0]
    request_id = str(uuid4())
    body = {"id": request_id, "expected_version": 0, "decision": "addressed", "reason": "已核对 " + "9" * 32}
    assert client.post(f'{url}/{row["id"]}/handling', json=body).status_code == 200
    other_record = record(client, c["id"], "SYN-OTHER", client.get(f'/api/company-workspaces/{c["id"]}').json()["version"])
    assert client.get(url, params={"request_id": request_id, "record_id": other_record["id"]}).status_code == 404
    client.__exit__(None, None, None)
    with TestClient(create_desktop_app(data_dir=tmp_path, launch_token=HEADERS["X-Qian-Desktop-Token"])) as reopened:
        reopened.headers.update(HEADERS)
        state = reopened.get(url, params={"request_id": request_id}).json()
        assert state["request_handling"]["id"] == request_id
        assert "9" * 32 not in json.dumps(state)
        assert state["observations"][0]["source"] == row["source"]
        assert reopened.post(f'{url}/{row["id"]}/handling', json=body).status_code == 409


def test_ambiguous_candidate_alternatives_have_explicit_advisory_matching_scope(api, tmp_path):
    from sqlalchemy import select
    from qian_labor.ai.grounding import ground_advisory
    from qian_labor.models.core import ContractClauseObservation, UploadedFile, EmployeeMatchCandidate
    from qian_labor.services.contract_advisory import persist_run
    client, db, c, rec, aid, provider, pipeline = processed_contract(api, tmp_path)
    result = ExtractionResult(contract_advisory=advisory(), employee_number="SYN-001")
    grounded = ground_advisory(result.contract_advisory, contract_input(), "contract.docx")
    with db.session() as s:
        file = s.scalar(select(UploadedFile))
        alternatives = [EmployeeMatchCandidate(analysis_id=aid, file_id=file.id, score=0.5, reason="synthetic_ambiguous",
            extracted_fields={"employee_ids": ["SYN-001"], "fact_ids": []}) for _ in range(2)]
        s.add_all(alternatives)
        s.flush()
        run = persist_run(s, aid, file, "synthetic-explicit-scope", [(result, None, [])], [grounded], [(None, alternatives)])
        s.flush()
        rows = list(s.scalars(select(ContractClauseObservation).where(ContractClauseObservation.run_id == run.id)))
        assert all(row.employee_id is None and row.match_candidate_id is not None for row in rows)
        assert rows[0].match_candidate_id not in {a.id for a in alternatives}


def test_actual_historical_adoption_keeps_advisory_get_readonly_and_rejects_handling(api, tmp_path):
    from uuid import uuid4
    from sqlalchemy import select
    from qian_labor.models.core import Employee, ContractClauseObservation, ContractAdvisoryRun, AIUsageRecord
    from qian_labor.desktop.import_service import DesktopImportService
    client, db = api
    c = company(client)
    aid = client.post('/api/analyses', json={"name": "合成历史", "company_display_name": "合成企业"}).json()["id"]
    doc = Document()
    doc.add_paragraph(advisory()["observations"][0]["source"]["excerpt"])
    path = tmp_path / "history-contract.docx"
    doc.save(path)
    imports = DesktopImportService(db, tmp_path)
    imports.import_paths(aid, [path])
    provider = SyntheticClauseProvider()
    ProcessingPipeline(db, imports.storage, provider=provider).process(aid)
    with db.session() as s:
        employee = s.scalar(select(Employee).where(Employee.analysis_id == aid))
        employee_id = employee.id
    response = client.put(f'/api/company-workspaces/{c["id"]}/analyses/{aid}/binding', json={
        "expected_company_version": 0, "expected_analysis_version": 0, "decisions": [
            {"action": "create", "snapshot_id": employee_id, "record": {"id": str(uuid4()), "display_name": "合成历史员工", "employee_number": "SYN-001"}}]})
    assert response.status_code == 200, response.text
    def snapshot():
        with db.session() as s:
            return [(r.id, r.source_location, r.source_excerpt, r.version) for r in s.scalars(select(ContractClauseObservation))], [r.id for r in s.scalars(select(AIUsageRecord))]
    before = snapshot()
    url = f'/api/company-workspaces/{c["id"]}/analyses/{aid}/contract-advisories'
    state = client.get(url).json()
    assert state["read_only"] and state["observations"][0]["read_only"]
    response = client.post(f'{url}/{state["observations"][0]["id"]}/handling', json={"id": str(uuid4()), "expected_version": 0, "decision": "checking", "reason": "合成核对"})
    assert response.status_code == 409 and response.json()["detail"]["code"] == "WORKSPACE_HISTORICAL_READ_ONLY"
    assert snapshot() == before and provider.calls == 1


def test_real_two_page_parser_aggregates_partial_advisory_without_promoting_one_page(api, tmp_path):
    from synthetic_mixed_materials import build_materials
    from qian_labor.desktop.import_service import DesktopImportService
    client, db = api
    c = company(client)
    aid = current(client, c["id"])["analysis_id"]
    path = tmp_path / "renewal.pdf"
    path.write_bytes(build_materials()[path.name])
    imports = DesktopImportService(db, tmp_path)
    imports.import_paths(aid, [path])
    class MultiPage(SyntheticClauseProvider):
        def extract(self, filename, content):
            self.calls += 1
            payload = advisory("SYN-001 renewal notice 2027-01-01", page=1) if self.calls == 1 else {
                "version": "contract-advisory-v1", "status": "unreadable", "observations": []}
            return ExtractionResult(document_type="contract", contract_advisory=payload)
    provider = MultiPage()
    ProcessingPipeline(db, imports.storage, provider=provider).process(aid)
    state = client.get(f'/api/company-workspaces/{c["id"]}/analyses/{aid}/contract-advisories').json()
    assert provider.calls == 2
    assert state["runs"][0]["execution_status"] == "partial"
    assert state["runs"][0]["input_count"] == 2 and state["runs"][0]["completed_input_count"] == 1
    assert len(state["observations"]) == 1 and state["observations"][0]["source"]["location"]["page"] == 1


@pytest.mark.parametrize("already_bound", [False, True])
def test_mixed_advisory_identity_cannot_steal_canonical_top_level_fact_scope(api, tmp_path, already_bound):
    from sqlalchemy import select
    from qian_labor.models.core import EmploymentFact, UploadedFile, EmployeeSnapshotBinding
    from qian_labor.desktop.import_service import DesktopImportService
    client, db, c, rec_a, aid, _, _ = processed_contract(api, tmp_path)
    if already_bound:
        seed = client.get(f'/api/analyses/{aid}/matching-candidates').json()["candidates"][0]
        assert select_record(client, aid, seed, rec_a).status_code == 200
    rec_b = record(client, c["id"], "SYN-002", client.get(f'/api/company-workspaces/{c["id"]}').json()["version"])
    document = Document()
    document.add_paragraph("SYN-001 signed contract")
    document.add_paragraph("SYN-002 wage clause requires checking")
    path = tmp_path / "mixed-identities.docx"
    document.save(path)
    imports = DesktopImportService(db, tmp_path)
    imports.import_paths(aid, [path])
    class MixedProvider(SyntheticClauseProvider):
        def extract(self, filename, content):
            payload = advisory("SYN-002 wage clause requires checking")
            payload["observations"][0]["employee_number"] = "SYN-002"
            return ExtractionResult(employee_number="SYN-001", document_type="contract", contract_advisory=payload,
                facts=[{"employee_id": None, "fact_type": "employment.contract.exists", "value": True,
                        "confidence": 1, "source": {"file_name": filename, "excerpt": "SYN-001 signed contract"}}])
    ProcessingPipeline(db, imports.storage, provider=MixedProvider()).process(aid)
    with db.session() as s:
        file = s.scalar(select(UploadedFile).where(UploadedFile.original_filename == path.name))
        file_id = file.id
        fact = s.scalar(select(EmploymentFact).where(EmploymentFact.file_id == file_id))
        fact_id = fact.id
        if already_bound:
            binding = s.scalar(select(EmployeeSnapshotBinding).where(EmployeeSnapshotBinding.employee_record_id == rec_a["id"]))
            assert fact.employee_id == binding.snapshot_id, "Advisory SYN-002 must not unset an owned canonical SYN-001 fact"
    candidates = [x for x in client.get(f'/api/analyses/{aid}/matching-candidates').json()["candidates"] if x["file_id"] == file_id]
    b = next(x for x in candidates if x["extracted_fields"]["employee_ids"] == ["SYN-002"])
    assert b["fact_ids"] == []
    if not already_bound:
        a = next(x for x in candidates if x["extracted_fields"]["employee_ids"] == ["SYN-001"])
        assert a["fact_ids"] == [fact_id], "Top-level canonical identity must retain a reachable matching scope"
        assert select_record(client, aid, a, rec_a).status_code == 200
    assert select_record(client, aid, b, rec_b).status_code == 200
    with db.session() as s:
        a_binding = s.scalar(select(EmployeeSnapshotBinding).where(EmployeeSnapshotBinding.employee_record_id == rec_a["id"]))
        assert s.get(EmploymentFact, fact_id).employee_id == a_binding.snapshot_id
    rows = client.get(f'/api/company-workspaces/{c["id"]}/analyses/{aid}/contract-advisories', params={"file_id": file_id, "record_id": rec_b["id"]}).json()["observations"]
    assert len(rows) == 1 and rows[0]["source"]["excerpt"].startswith("SYN-002 ")
