"""Synthetic parser-grounding tests; OCR is injected only at the OCR boundary."""
from datetime import datetime
from io import BytesIO

import pytest
from docx import Document
from openpyxl import Workbook

from qian_labor.ai.schemas import ExtractionResult
from qian_labor.jobs.processing import ProcessingPipeline
from qian_labor.parsers.registry import ParserRegistry


def word_bytes():
    doc = Document()
    doc.add_paragraph("SYN-001 signed contract")
    doc.add_paragraph("SYN-001 signed contract")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "SYN-001"
    table.cell(0, 1).text = "Renewal 2026-09-01"
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def result(excerpt, **position):
    return ExtractionResult.model_validate({"employee_number": "SYN-001", "facts": [{
        "employee_id": "SYN-001", "fact_type": "employment.contract.exists", "value": True,
        "confidence": 1, "source": {"file_name": "invented-page-99.txt", "excerpt": excerpt, **position},
    }]})


def test_docx_inputs_keep_parser_context_and_hints():
    content = word_bytes()
    parsed = ParserRegistry().parse("contract.docx", content)
    item = ProcessingPipeline._extraction_inputs("contract.docx", content, parsed)[0]
    assert hasattr(item, "blocks"), "Extraction input drops authoritative parser context"
    assert item.blocks == tuple(parsed.blocks)
    assert b'"paragraph": 1' in item.content
    assert b'"table": 1' in item.content


def test_extraction_request_carries_stable_source_ids_and_grounds_by_id():
    import hashlib
    from qian_labor.ai.grounding import ground_result
    from qian_labor.services.source_provenance import deterministic_citation_id

    workbook = Workbook()
    workbook.active.title = "工资表"
    workbook.active.append(["员工编号", "实际工资"])
    workbook.active.append(["SYN-001", 9000])
    output = BytesIO()
    workbook.save(output)
    content = output.getvalue()
    parsed = ParserRegistry().parse("工资表.xlsx", content)
    item = ProcessingPipeline._extraction_inputs("工资表.xlsx", content, parsed)[0]
    wage = next(block for block in parsed.blocks if block.locator.get("cell") == "B2")
    expected = deterministic_citation_id(
        hashlib.sha256(content).hexdigest(),
        {key: wage.locator[key] for key in ("sheet", "row", "column", "cell")},
        wage.text,
    )

    assert expected in item.content.decode("utf-8")
    extracted = ExtractionResult.model_validate({
        "employee_number": "SYN-001",
        "facts": [{
            "employee_id": "SYN-001",
            "fact_type": "employment.pay.actual_wage",
            "value": 9000,
            "confidence": 1,
            "source": {
                "file_name": "model-output",
                "excerpt": wage.text,
                "sheet": "工资表",
                "row": 2,
                "column": "2",
                "cell": "B2",
                "citation_id": expected,
            },
        }],
    })

    proofs = ground_result(extracted, item, "工资表.xlsx")

    assert proofs[0]["status"] == "locally_located"
    assert extracted.facts[0].source.sheet == "工资表"
    assert extracted.facts[0].source.row == 2
    assert extracted.facts[0].source.cell == "B2"
    assert extracted.facts[0].source.excerpt == wage.text
    assert expected == deterministic_citation_id(
        hashlib.sha256(content).hexdigest(),
        extracted.facts[0].source.model_dump(
            exclude={"file_name", "excerpt"}, exclude_none=True
        ),
        extracted.facts[0].source.excerpt,
    )

    row_text = "SYN-001 | 9000"
    row_citation_id = deterministic_citation_id(
        hashlib.sha256(content).hexdigest(),
        {"sheet": "工资表", "row": 2},
        row_text,
    )
    assert row_citation_id in item.content.decode("utf-8")
    row_extracted = ExtractionResult.model_validate({
        "employee_number": "SYN-001",
        "facts": [{
            "employee_id": "SYN-001",
            "fact_type": "employment.pay.actual_wage",
            "value": 9000,
            "confidence": 1,
            "source": {
                "file_name": "model-output",
                "excerpt": row_text,
                "sheet": "工资表",
                "row": 2,
                "citation_id": row_citation_id,
            },
        }],
    })

    row_proofs = ground_result(row_extracted, item, "工资表.xlsx")

    assert row_proofs[0]["status"] == "locally_located"
    assert row_proofs[0]["requires_review"] is True
    assert row_extracted.facts[0].needs_human_confirmation is True
    assert row_extracted.facts[0].source.row == 2
    assert row_extracted.facts[0].source.column is None
    assert row_extracted.facts[0].source.excerpt == row_text


def test_cited_empty_or_partial_excerpt_requires_human_confirmation():
    import hashlib
    from qian_labor.ai.grounding import ExtractionInput, deterministic_citation_id, ground_result
    from qian_labor.ai.schemas import EmploymentFact, SourceLocator
    from qian_labor.parsers.protocols import ParsedBlock

    content = b"synthetic spreadsheet bytes"
    blocks = (
        ParsedBlock("SYN-001", "cell", {"sheet": "工资表", "row": 2, "column": 1, "cell": "A2"}),
        ParsedBlock("9000", "cell", {"sheet": "工资表", "row": 2, "column": 2, "cell": "B2"}),
    )
    item = ExtractionInput(
        "工资表.xlsx",
        content,
        blocks,
        source_file_hash=hashlib.sha256(content).hexdigest(),
    )
    citation_id = deterministic_citation_id(
        item.source_file_hash,
        blocks[1].locator,
        blocks[1].text,
    )

    for excerpt in ("", "900"):
        extracted = ExtractionResult(facts=[EmploymentFact(
            employee_id="SYN-001",
            fact_type="employment.pay.actual_wage",
            value=9000,
            confidence=1,
            source=SourceLocator(
                file_name="model-output",
                sheet="工资表",
                row=2,
                column="2",
                cell="B2",
                excerpt=excerpt,
                citation_id=citation_id,
            ),
        )])

        proofs = ground_result(extracted, item, "工资表.xlsx")

        assert proofs[0]["status"] == "locally_located"
        assert proofs[0]["requires_review"] is True
        assert extracted.facts[0].needs_human_confirmation is True
        assert extracted.facts[0].source.cell == "B2"
        assert extracted.facts[0].source.excerpt == "9000"


def test_fabricated_source_id_fails_closed_even_with_an_exact_quote():
    from qian_labor.ai.grounding import ground_result

    workbook = Workbook()
    workbook.active.append(["员工编号", "实际工资"])
    workbook.active.append(["SYN-001", 9000])
    output = BytesIO()
    workbook.save(output)
    content = output.getvalue()
    parsed = ParserRegistry().parse("工资表.xlsx", content)
    item = ProcessingPipeline._extraction_inputs("工资表.xlsx", content, parsed)[0]
    extracted = ExtractionResult.model_validate({
        "employee_number": "SYN-001",
        "facts": [{
            "employee_id": "SYN-001",
            "fact_type": "employment.pay.actual_wage",
            "value": 9000,
            "confidence": 1,
            "source": {
                "file_name": "model-output",
                "excerpt": "9000",
                "sheet": "Sheet",
                "row": 2,
                "column": "2",
                "cell": "B2",
                "citation_id": "cite-fabricated-source-id",
            },
        }],
    })

    proofs = ground_result(extracted, item, "工资表.xlsx")

    assert proofs[0]["status"] == "unlocated_needs_review"
    assert extracted.facts[0].source.excerpt == ""
    assert extracted.facts[0].needs_human_confirmation


def test_grounding_preserves_all_actual_occurrences_and_table_identity():
    from qian_labor.ai.grounding import ground_result
    content = word_bytes()
    item = ProcessingPipeline._extraction_inputs("contract.docx", content, ParserRegistry().parse("contract.docx", content))[0]
    extracted = result("SYN-001  signed\ncontract")
    proofs = ground_result(extracted, item, "contract.docx")
    assert [f.source.paragraph for f in extracted.facts] == [1, 2]
    assert all(p["status"] == "locally_located" for p in proofs)
    assert all(f.needs_human_confirmation for f in extracted.facts)
    extracted = result("Renewal 2026-09-01")
    ground_result(extracted, item, "contract.docx")
    assert extracted.facts[0].source.table == 1
    assert extracted.facts[0].source.row == 1
    assert extracted.facts[0].source.column == "2"
    assert extracted.facts[0].source.sheet is None


@pytest.mark.parametrize("quote,position", [("invented signed excerpt", {}), ("SYN-001 signed contract", {"page": 99}), ("SYN-001 signed contract", {"paragraph": 99})])
def test_fabricated_quote_or_coordinate_is_not_original_evidence(quote, position):
    from qian_labor.ai.grounding import ground_result
    content = word_bytes()
    item = ProcessingPipeline._extraction_inputs("contract.docx", content, ParserRegistry().parse("contract.docx", content))[0]
    extracted = result(quote, **position)
    proofs = ground_result(extracted, item, "contract.docx")
    assert proofs[0]["status"] == "unlocated_needs_review"
    assert extracted.facts[0].source.excerpt == ""
    assert extracted.facts[0].source.page is None
    assert extracted.facts[0].needs_human_confirmation


def test_xlsx_actual_dates_are_readable_dates():
    workbook = Workbook()
    workbook.active.append(["employee", "start"])
    workbook.active.append(["SYN-001", datetime(2026, 9, 1)])
    output = BytesIO()
    workbook.save(output)
    parsed = ParserRegistry().parse("dates.xlsx", output.getvalue())
    date = next(b for b in parsed.blocks if b.locator.get("cell") == "B2")
    assert date.text == "2026-09-01"


def test_extraction_key_is_versioned_but_parse_key_is_not():
    from qian_labor.ai.grounding import EXTRACTION_VERSION
    key = ProcessingPipeline._job_key("a", "f", "extract", "digest")
    assert key != "a:f:extract:digest"
    assert EXTRACTION_VERSION in key
    assert ProcessingPipeline._job_key("a", "f", "parse", "digest") == "a:f:parse:digest"


def test_v3_extraction_does_not_reuse_v1_key_but_keeps_parse_key():
    from qian_labor.ai.grounding import EXTRACTION_VERSION

    assert EXTRACTION_VERSION == "parser-grounding-v3"
    assert ProcessingPipeline._job_key("a", "f", "parse", "digest") == "a:f:parse:digest"
    assert ProcessingPipeline._job_key("a", "f", "extract", "digest") != (
        "a:f:extract:digest:parser-grounding-v1:contract-advisory-v1"
    )


def test_source_context_upgrade_does_not_reuse_previous_v2_job_key():
    old_key = "a:f:extract:digest:parser-grounding-v2:contract-advisory-v1"
    assert ProcessingPipeline._job_key("a", "f", "extract", "digest") != old_key
    assert ProcessingPipeline._job_key("a", "f", "parse", "digest") == "a:f:parse:digest"


def test_unlocated_spreadsheet_sources_disable_cache_until_explicit_retry(tmp_path):
    import hashlib
    from fastapi.testclient import TestClient
    from sqlalchemy import select
    from qian_labor.ai.grounding import EXTRACTION_VERSION
    from qian_labor.desktop.app import create_desktop_app
    from qian_labor.models.core import AnalysisBatch, EmploymentFact, ProcessingJob, SourceLocator, UploadedFile
    from test_workspace_api import HEADERS, TOKEN

    class Provider:
        name, is_external = "fake", False

    app = create_desktop_app(data_dir=tmp_path / "app", launch_token=TOKEN)
    with TestClient(app) as client:
        aid = client.post("/api/analyses", headers=HEADERS, json={
            "name": "spreadsheet source retry", "company_display_name": "fictional",
        }).json()["id"]
        workbook_path = tmp_path / "records.xlsx"
        workbook = Workbook()
        workbook.active.append(["员工编号", "试用期开始"])
        workbook.active.append(["SYN-001", "2026-09-01"])
        output = BytesIO()
        workbook.save(output)
        content = output.getvalue()
        workbook_path.write_bytes(content)
        imported = client.post(f"/api/analyses/{aid}/import-paths", headers=HEADERS,
                               json={"paths": [str(workbook_path)]})
        imported.raise_for_status()
        with app.state.database.session() as session:
            file = session.scalar(select(UploadedFile).where(UploadedFile.analysis_id == aid))
            assert file is not None
            file.status = "processed"
            file.detected_kind = "spreadsheet"
            job_key = ProcessingPipeline._job_key(aid, file.id, "extract", file.sha256)
            session.add(ProcessingJob(analysis_id=aid, file_id=file.id, job_type="extract",
                input_hash=file.sha256, unique_key=job_key, status="succeeded", attempts=1))
            fact = EmploymentFact(analysis_id=aid, file_id=file.id, fact_type="employment.probation.start",
                value_json="2026-09-01", normalized_value_json="2026-09-01", extraction_method="fake",
                confidence=1.0, verification_status="needs_human_confirmation",
                dedupe_key=hashlib.sha256(b"synthetic-unlocated-fact").hexdigest())
            session.add(fact)
            session.flush()
            session.add(SourceLocator(analysis_id=aid, file_id=file.id, fact_id=fact.id,
                locator_type="document", location={"_grounding": {
                    "version": EXTRACTION_VERSION, "status": "unlocated_needs_review",
                }}, excerpt="", content_hash=hashlib.sha256(b"").hexdigest()))
            session.get(AnalysisBatch, aid).status = "completed"
            session.commit()
            assert ProcessingPipeline.has_unlocated_sources(session, aid, file.id) is True
            assert ProcessingPipeline.cached_extraction(session, aid, file, Provider()) is False
        workspace = client.get(f"/api/analyses/{aid}/workspace", headers=HEADERS)
        workspace.raise_for_status()
        assert workspace.json()["files"][0]["needs_reextraction"] is True


def test_previous_v2_success_cache_requires_explicit_source_context_upgrade(tmp_path):
    from fastapi.testclient import TestClient
    from sqlalchemy import select
    from qian_labor.desktop.app import create_desktop_app
    from qian_labor.models.core import AnalysisBatch, ProcessingJob, UploadedFile
    from qian_labor.storage.local import LocalStorage
    from test_workspace_api import HEADERS, TOKEN

    class SyntheticProvider:
        name, is_external = "fake", False

        def extract(self, filename, content):
            return result("SYN-001 signed contract")

    app = create_desktop_app(data_dir=tmp_path / "app", launch_token=TOKEN)
    with TestClient(app) as client:
        response = client.post("/api/analyses", headers=HEADERS, json={
            "name": "synthetic cache upgrade", "company_display_name": "fictional",
        })
        response.raise_for_status()
        aid = response.json()["id"]
        path = tmp_path / "contract.docx"
        path.write_bytes(word_bytes())
        imported = client.post(f"/api/analyses/{aid}/import-paths", headers=HEADERS,
                               json={"paths": [str(path)]})
        imported.raise_for_status()
        provider = SyntheticProvider()
        pipeline = ProcessingPipeline(app.state.database,
            LocalStorage(str(app.state.storage_root)), provider)
        pipeline.process(aid)
        with app.state.database.session() as session:
            file = session.scalar(select(UploadedFile).where(UploadedFile.analysis_id == aid))
            job = session.scalar(select(ProcessingJob).where(
                ProcessingJob.analysis_id == aid, ProcessingJob.job_type == "extract"))
            assert file is not None and job is not None and job.status == "succeeded"
            job.unique_key = f"{aid}:{file.id}:extract:{file.sha256}:parser-grounding-v2:contract-advisory-v1"
            session.get(AnalysisBatch, aid).status = "completed"
            session.commit()
            assert ProcessingPipeline.cached_extraction(session, aid, file, provider) is False
        workspace = client.get(f"/api/analyses/{aid}/workspace", headers=HEADERS)
        workspace.raise_for_status()
        assert workspace.json()["files"][0]["needs_reextraction"] is True


def test_v2_retry_runs_once_then_v3_cache_reuses_without_read_calls(tmp_path):
    from fastapi.testclient import TestClient
    from sqlalchemy import select
    from qian_labor.desktop.app import create_desktop_app
    from qian_labor.models.core import AnalysisBatch, ProcessingJob, UploadedFile
    from qian_labor.storage.local import LocalStorage
    from test_workspace_api import HEADERS, TOKEN

    class CountingProvider:
        name, is_external, calls = "fake", False, 0

        def extract(self, filename, content):
            self.calls += 1
            return result("SYN-001 signed contract")

    app = create_desktop_app(data_dir=tmp_path / "app", launch_token=TOKEN)
    with TestClient(app) as client:
        response = client.post("/api/analyses", headers=HEADERS, json={
            "name": "synthetic retry", "company_display_name": "fictional",
        })
        response.raise_for_status()
        aid = response.json()["id"]
        path = tmp_path / "contract.docx"
        path.write_bytes(word_bytes())
        imported = client.post(f"/api/analyses/{aid}/import-paths", headers=HEADERS,
                               json={"paths": [str(path)]})
        imported.raise_for_status()
        provider = CountingProvider()
        pipeline = ProcessingPipeline(app.state.database,
            LocalStorage(str(app.state.storage_root)), provider)
        with app.state.database.session() as session:
            file = session.scalar(select(UploadedFile).where(UploadedFile.analysis_id == aid))
            assert file is not None
            file_id, file_hash = file.id, file.sha256
        pipeline._process_file(aid, file_id)
        first_count = provider.calls
        assert first_count == 1

        with app.state.database.session() as session:
            job = session.scalar(select(ProcessingJob).where(
                ProcessingJob.analysis_id == aid, ProcessingJob.file_id == file_id,
                ProcessingJob.job_type == "extract"))
            assert job is not None and job.status == "succeeded"
            job.unique_key = f"{aid}:{file_id}:extract:{file_hash}:parser-grounding-v2:contract-advisory-v1"
            session.get(AnalysisBatch, aid).status = "created"
            session.commit()

        pipeline._process_file(aid, file_id)
        assert provider.calls == first_count + 1
        pipeline._process_file(aid, file_id)
        assert provider.calls == first_count + 1

        for url in (f"/api/analyses/{aid}/workspace", f"/api/analyses/{aid}/report"):
            response = client.get(url, headers=HEADERS)
            assert response.status_code == 200, response.text
        assert provider.calls == first_count + 1


def test_zhipu_text_limit_fails_before_transport():
    import httpx
    from qian_labor.ai.providers import AIProviderError
    from qian_labor.ai.zhipu_provider import ZhipuChatCompletionsProvider
    from qian_labor.security.local_redaction import PreparedProviderContent
    from qian_labor.security.local_redaction import PrivacyBoundary
    calls = []
    provider = ZhipuChatCompletionsProvider(api_key="synthetic", base_url="https://open.bigmodel.cn/api/paas/v4",
        text_model="glm-5.3-flash", vision_model="glm-5.3-flash",
        privacy_boundary=PrivacyBoundary("synthetic-grounding-pepper-at-least-32-characters"), max_attempts=1,
        client=httpx.Client(transport=httpx.MockTransport(lambda r: calls.append(r) or httpx.Response(500))))
    with pytest.raises(AIProviderError, match="AI_TEXT_LIMIT_EXCEEDED"):
        provider.extract("synthetic.txt", PreparedProviderContent(("x" * 100_001).encode()))
    assert calls == []


def test_required_ocr_context_is_reused_once_and_masked():
    from PIL import Image
    from qian_labor.security.local_redaction import PrivacyBoundary, LocalImageRedactor, OCRToken
    image = BytesIO()
    Image.new("RGB", (400, 80), "white").save(image, format="PNG")
    class SyntheticOCR:
        calls = 0
        def extract_tokens(self, content):
            self.calls += 1
            return [OCRToken("SYN-001 signed contract", 2, 4, 200, 20, "line1")]
    ocr = SyntheticOCR()
    prepared = PrivacyBoundary("synthetic-pepper", LocalImageRedactor(ocr)).prepare("scan.png", image.getvalue(), is_image=True, external=True)
    assert ocr.calls == 1
    assert hasattr(prepared, "ocr_blocks"), "Required redaction OCR context is discarded"
    assert prepared.ocr_blocks[0].text == "SYN-001 signed contract"
    assert prepared.ocr_blocks[0].locator["bbox"] == [2, 4, 202, 24]


def test_vision_request_context_keeps_embedded_image_locator_and_ocr_identity():
    import hashlib
    from qian_labor.ai.grounding import ExtractionInput, deterministic_citation_id
    from qian_labor.parsers.protocols import ParsedBlock
    from qian_labor.security.local_redaction import PreparedProviderInput

    image_bytes = b"synthetic-image"
    item = ExtractionInput(
        "embedded-warning.docx-image-1.png",
        image_bytes,
        (ParsedBlock("", "image_context", {"paragraph": 2, "image": 1}),),
        None,
        hashlib.sha256(image_bytes).hexdigest(),
    )
    prepared = PreparedProviderInput(
        item.filename,
        image_bytes,
        {},
        ocr_blocks=(ParsedBlock(
            "SYN-001 signed contract",
            "ocr_line",
            {"block": 1, "bbox": [2, 4, 202, 24]},
        ),),
    )

    context = ProcessingPipeline._vision_source_context(item, prepared)

    assert context["parser_context"] == [{"locator": {"paragraph": 2, "image": 1}}]
    ocr = context["ocr_blocks"][0]
    assert ocr["locator"] == {
        "paragraph": 2, "image": 1, "block": 1, "bbox": [2, 4, 202, 24]
    }
    assert ocr["citation_id"] == deterministic_citation_id(
        item.source_file_hash,
        ocr["locator"],
        "SYN-001 signed contract",
    )


def test_legacy_sources_are_readonly_labeled_and_new_unlocated_forces_review(review_case):
    from sqlalchemy import select
    from test_finding_review_api import HEADERS, stored_state
    from qian_labor.models.core import SourceLocator, RiskFinding
    client, database, analysis_id, _, finding_id = review_case
    before = stored_state(database, finding_id)
    with database.session() as session:
        original_sources = [(s.id, s.location, s.excerpt) for s in session.scalars(select(SourceLocator).order_by(SourceLocator.id))]
    detail = client.get(f"/api/findings/{finding_id}", headers=HEADERS).json()
    assert all(s.get("provenance") == "legacy_unverified" for s in detail["sources"])
    assert stored_state(database, finding_id) == before
    with database.session() as session:
        assert [(s.id, s.location, s.excerpt) for s in session.scalars(select(SourceLocator).order_by(SourceLocator.id))] == original_sources
    with database.session() as session:
        finding = session.get(RiskFinding, finding_id)
        source = session.get(SourceLocator, finding.source_locator_ids[0])
        source.location = {"page": 999, "_grounding": {"version": "parser-grounding-v1", "status": "unlocated_needs_review"}}
        source.excerpt = "invented model excerpt"
        session.commit()
    detail = client.get(f"/api/findings/{finding_id}", headers=HEADERS).json()
    unlocated = next(s for s in detail["sources"] if s["provenance"] == "unlocated_needs_review")
    assert unlocated["excerpt"] == "" and unlocated["location"] == {}
    assert detail["requires_human_review"] is True
    report = client.get(f"/api/analyses/{analysis_id}/report", headers=HEADERS).json()
    item = next(f for f in report["findings"] if f["id"] == finding_id)
    assert any(s["provenance"] == "unlocated_needs_review" and s["location"] == {} for s in item["sources"])
    from test_finding_review_api import post_review
    from qian_labor.models.core import EmploymentFact
    with database.session() as session:
        originals = [(f.id, f.value_json, f.normalized_value_json, f.verification_status) for f in session.scalars(select(EmploymentFact).order_by(EmploymentFact.id))]
    reviewed = post_review(client, finding_id, status="reviewed")
    assert reviewed.status_code == 200
    assert any(s["provenance"] == "unlocated_needs_review" for s in reviewed.json()["sources"])
    with database.session() as session:
        assert [(f.id, f.value_json, f.normalized_value_json, f.verification_status) for f in session.scalars(select(EmploymentFact).order_by(EmploymentFact.id))] == originals


from test_finding_review_api import review_case  # shared synthetic database fixture


def test_mixed_fixture_actual_structure():
    from synthetic_mixed_materials import build_materials, EMPLOYEES
    materials = build_materials()
    parsed = {name: ParserRegistry().parse(name, content) for name, content in materials.items()}
    assert len(EMPLOYEES) == 10
    assert [b.locator["row"] for b in parsed["roster.csv"].blocks[1:]] == list(range(2, 12))
    assert [b.text.split()[0] for b in parsed["contracts.docx"].blocks[:10]] == list(EMPLOYEES)
    assert {b.locator["page"] for b in parsed["renewal.pdf"].blocks} == {1, 2}
    assert parsed["probation.xlsx"].blocks[2].text == "2026-03-01"
    assert parsed["termination-scan.pdf"].vision_pages[0].page == 1
    assert parsed["termination.png"].blocks == []
    assert parsed["embedded-warning.docx"].needs_vision
    assert len(parsed["embedded-warning.docx"].vision_pages) == 1


def test_partial_parser_warning_and_old_cache_upgrade_are_explicit(tmp_path):
    from fastapi.testclient import TestClient
    from sqlalchemy import select
    from synthetic_mixed_materials import build_materials
    from qian_labor.desktop.app import create_desktop_app
    from qian_labor.models.core import UploadedFile, ProcessingJob, AnalysisBatch, EmploymentFact, SourceLocator
    from qian_labor.storage.local import LocalStorage
    from test_workspace_api import HEADERS, TOKEN
    class GroundedProvider:
        name, is_external = "fake", False
        calls = 0
        def extract(self, filename, content):
            self.calls += 1
            return result("SYN-001 signed contract")
    provider = GroundedProvider()
    app = create_desktop_app(data_dir=tmp_path / "app", launch_token=TOKEN)
    with TestClient(app) as client:
        aid = client.post("/api/analyses", headers=HEADERS, json={"name": "synthetic", "company_display_name": "fictional"}).json()["id"]
        path = tmp_path / "embedded-warning.docx"
        path.write_bytes(build_materials()[path.name])
        client.post(f"/api/analyses/{aid}/import-paths", headers=HEADERS, json={"paths": [str(path)]})
        pipeline = ProcessingPipeline(app.state.database, LocalStorage(str(app.state.storage_root)), provider)
        with app.state.database.session() as session:
            file = session.scalar(select(UploadedFile))
            file.status = "processed"
            session.add(ProcessingJob(analysis_id=aid, file_id=file.id, job_type="extract", input_hash=file.sha256,
                unique_key=f"{aid}:{file.id}:extract:{file.sha256}", status="succeeded"))
            session.get(AnalysisBatch, aid).status = "completed"
            session.commit()
        workspace = client.get(f"/api/analyses/{aid}/workspace", headers=HEADERS).json()
        assert workspace["files"][0].get("needs_reextraction") is True
        pipeline.process(aid)  # Generic completed pipeline stays a no-op.
        assert provider.calls == 0
        app.state.processing_queue.pipeline_factory = lambda: pipeline
        assert client.post(f"/api/analyses/{aid}/process", headers=HEADERS).status_code == 202
        app.state.processing_queue.shutdown()
        workspace = client.get(f"/api/analyses/{aid}/workspace", headers=HEADERS).json()
        assert provider.calls == 2
        assert workspace["analysis"]["status"] in {"matching_review", "partial", "evaluating", "completed"}
        assert workspace["files"][0]["warnings"] == []
        assert workspace["files"][0]["fact_count"] == 2
        assert workspace["files"][0]["needs_reextraction"] is False
        with app.state.database.session() as session:
            assert len(list(session.scalars(select(ProcessingJob).where(ProcessingJob.job_type == "extract")))) == 2
        client.get(f"/api/analyses/{aid}/workspace", headers=HEADERS)
        assert provider.calls == 2


def test_legacy_v1_success_job_is_not_reusable_after_grounding_upgrade(tmp_path):
    from fastapi.testclient import TestClient
    from sqlalchemy import select
    from qian_labor.desktop.app import create_desktop_app
    from qian_labor.models.core import AnalysisBatch, ProcessingJob, UploadedFile
    from qian_labor.storage.local import LocalStorage
    from test_workspace_api import HEADERS, TOKEN

    class GroundedProvider:
        name, is_external = "fake", False

        def extract(self, filename, content):
            return result("SYN-001 signed contract")

    app = create_desktop_app(data_dir=tmp_path / "app", launch_token=TOKEN)
    with TestClient(app) as client:
        aid = client.post("/api/analyses", headers=HEADERS,
                          json={"name": "legacy cache", "company_display_name": "fictional"}).json()["id"]
        path = tmp_path / "contract.docx"
        path.write_bytes(word_bytes())
        client.post(f"/api/analyses/{aid}/import-paths", headers=HEADERS,
                    json={"paths": [str(path)]})
        pipeline = ProcessingPipeline(app.state.database, LocalStorage(str(app.state.storage_root)), GroundedProvider())
        pipeline.process(aid)
        with app.state.database.session() as session:
            file = session.scalar(select(UploadedFile).where(UploadedFile.analysis_id == aid))
            job = session.scalar(select(ProcessingJob).where(
                ProcessingJob.analysis_id == aid, ProcessingJob.job_type == "extract"))
            assert file is not None and job is not None and job.status == "succeeded"
            job.unique_key = f"{aid}:{file.id}:extract:{file.sha256}:parser-grounding-v1:contract-advisory-v1"
            session.get(AnalysisBatch, aid).status = "completed"
            session.commit()
            assert ProcessingPipeline.cached_extraction(session, aid, file, GroundedProvider()) is False
        workspace = client.get(f"/api/analyses/{aid}/workspace", headers=HEADERS).json()
        assert workspace["files"][0]["extraction_version"] is None
        assert workspace["files"][0]["needs_reextraction"] is True


def test_v1_fact_and_revision_are_preserved_but_current_projection_uses_v2(tmp_path):
    import hashlib
    from sqlalchemy import select
    from qian_labor.models.core import (
        AnalysisBatch, CompanyAnalysisBinding, CompanyWorkspace, EmploymentFact,
        SourceLocator, UploadedFile,
    )
    from qian_labor.services.effective_facts import effective_projection
    from qian_labor.services.source_provenance import deterministic_citation_id
    from fastapi.testclient import TestClient
    from qian_labor.desktop.app import create_desktop_app
    from qian_labor.storage.local import LocalStorage
    from test_workspace_api import HEADERS, TOKEN

    class GroundedProvider:
        name, is_external = "fake", False

        def extract(self, filename, content):
            return result("SYN-001 signed contract")

    app = create_desktop_app(data_dir=tmp_path / "app", launch_token=TOKEN)
    with TestClient(app) as client:
        aid = client.post("/api/analyses", headers=HEADERS,
                          json={"name": "versioned facts", "company_display_name": "fictional"}).json()["id"]
        path = tmp_path / "contract.docx"
        path.write_bytes(word_bytes())
        client.post(f"/api/analyses/{aid}/import-paths", headers=HEADERS,
                    json={"paths": [str(path)]})
        with app.state.database.session() as session:
            file = session.scalar(select(UploadedFile).where(UploadedFile.analysis_id == aid))
            company = CompanyWorkspace(id="synthetic-version-company", display_name="fictional")
            session.add(company)
            session.flush()
            session.add(CompanyAnalysisBinding(analysis_id=aid, company_id=company.id, role="current"))
            old = EmploymentFact(
                analysis_id=aid, employee_id=None, file_id=file.id,
                fact_type="employment.contract.exists", value_json=True,
                normalized_value_json=True, extraction_method="synthetic-v1", confidence=1,
                verification_status="confirmed", dedupe_key="v1-" + aid,
            )
            session.add(old)
            session.flush()
            excerpt = "SYN-001 signed contract"
            location = {"paragraph": 1, "_grounding": {
                "version": "parser-grounding-v1", "status": "locally_located", "requires_review": False,
            }}
            location["_citation_id"] = deterministic_citation_id(file.sha256, location, excerpt)
            session.add(SourceLocator(
                analysis_id=aid, file_id=file.id, fact_id=old.id,
                locator_type="paragraph", location=location, excerpt=excerpt,
                content_hash=hashlib.sha256(excerpt.encode()).hexdigest(),
            ))
            session.commit()
            old_id = old.id
        pipeline = ProcessingPipeline(app.state.database, LocalStorage(str(app.state.storage_root)), GroundedProvider())
        pipeline.process(aid)
        with app.state.database.session() as session:
            facts = list(session.scalars(select(EmploymentFact).where(EmploymentFact.analysis_id == aid)))
            current_rows = effective_projection(session, aid)
            assert old_id in {fact.id for fact in facts}
            assert old_id not in {row.id for row in current_rows}
            assert any(row.id != old_id for row in current_rows)
            old_source = session.scalar(select(SourceLocator).where(SourceLocator.fact_id == old_id))
            assert old_source.location["_grounding"]["version"] == "parser-grounding-v1"
            from qian_labor.services.report_versions import ReportVersionService
            company_id = session.scalar(select(CompanyAnalysisBinding.company_id).where(
                CompanyAnalysisBinding.analysis_id == aid))
            payload, context = ReportVersionService(app.state.database, None).capture(session, company_id, aid)
            assert payload["facts"]
            assert context["input_revision"]
            session.scalar(select(CompanyAnalysisBinding).where(
                CompanyAnalysisBinding.analysis_id == aid)).role = "historical"
            historical_rows = effective_projection(session, aid)
            assert old_id in {row.id for row in historical_rows}
            assert next(row for row in historical_rows if row.id == old_id).state.valid is True
            assert session.get(AnalysisBatch, aid) is not None


def test_old_confirmed_fact_cannot_bypass_new_unlocated_uncertainty(review_case):
    from sqlalchemy import select
    from qian_labor.models.core import EmploymentFact, SourceLocator
    from qian_labor.services.risk_evaluation import RiskEvaluationService
    from qian_labor.services.assessment_scope import scoped_evidence_rows
    _, database, aid, _, _ = review_case
    with database.session() as session:
        fact = session.scalar(select(EmploymentFact))
        fact.verification_status = "confirmed"
        session.add(SourceLocator(analysis_id=aid, file_id=fact.file_id, fact_id=fact.id,
            locator_type="document", excerpt="", content_hash="synthetic", location={"_grounding": {
                "version": "parser-grounding-v1", "status": "unlocated_needs_review"}}))
        session.flush()
        assert RiskEvaluationService._context_facts(session, [fact])[fact.fact_type].conflicted
        assert all(r.verification_status == "needs_human_confirmation" for r in scoped_evidence_rows(session, aid) if r.fact_id == fact.id)
        assert fact.verification_status == "confirmed"


def test_unscoped_repeated_sentence_in_multi_employee_doc_is_not_attributed():
    from qian_labor.ai.grounding import ground_result
    doc = Document()
    for eid in ("SYN-001", "SYN-002"):
        doc.add_paragraph(eid)
        doc.add_paragraph("signed contract")
    output = BytesIO()
    doc.save(output)
    item = ProcessingPipeline._extraction_inputs("multi.docx", output.getvalue(), ParserRegistry().parse("multi.docx", output.getvalue()))[0]
    extracted = result("signed contract")
    # A single employee result still cannot assign the other employee's paragraph.
    assert ground_result(extracted, item, "multi.docx")[0]["status"] == "unlocated_needs_review"


def test_location_match_does_not_confirm_another_employee_interpretation():
    from qian_labor.ai.grounding import ground_result, ExtractionInput
    from qian_labor.parsers.protocols import ParsedBlock
    item = ExtractionInput("synthetic.docx", b"", (ParsedBlock("SYN-0010 signed contract", "paragraph", {"paragraph": 1}),))
    extracted = result("SYN-0010 signed contract")  # Provider attributes a different employee's paragraph to SYN-001.
    proof = ground_result(extracted, item, "synthetic.docx")
    assert extracted.facts[0].needs_human_confirmation is True
    assert proof[0]["requires_review"] is True


def test_pdf_block_number_and_filename_does_not_assert_page():
    from synthetic_mixed_materials import build_materials
    from qian_labor.ai.grounding import ground_result
    content = build_materials()["renewal.pdf"]
    parsed = ParserRegistry().parse("renewal.pdf", content)
    assert all("block" in b.locator for b in parsed.blocks)
    item = ProcessingPipeline._extraction_inputs("renewal.pdf", content, parsed)[1]
    extracted = result("SYN-006 renewal notice 2027-01-01")
    extracted.employee_number = extracted.facts[0].employee_id = "SYN-006"
    proofs = ground_result(extracted, item, "renewal.pdf")
    assert proofs[0]["status"] == "locally_located"
    assert extracted.facts[0].source.page == 2


@pytest.mark.parametrize("datemode,serial", [(0, 46266), (1, 44804)])
def test_xls_date_conversion_uses_workbook_epoch(monkeypatch, datemode, serial):
    from types import SimpleNamespace
    from unittest.mock import Mock
    import xlrd
    class Sheet:
        nrows, ncols, name = 2, 2, "synthetic"
        def cell_value(self, row, col):
            return [["employee", "start"], ["SYN-001", serial]][row][col]
        def cell_type(self, row, col):
            return xlrd.XL_CELL_DATE if (row, col) == (1, 1) else xlrd.XL_CELL_TEXT
    release_resources = Mock()
    monkeypatch.setattr(xlrd, "open_workbook", lambda **kwargs: SimpleNamespace(
        sheets=lambda: [Sheet()], datemode=datemode, release_resources=release_resources))
    parsed = ParserRegistry().parse("synthetic.xls", b"synthetic decoder boundary")
    assert parsed.blocks[-1].text == xlrd.xldate_as_datetime(serial, datemode).date().isoformat()
    assert parsed.blocks[-1].locator["cell"] == "B2"
    release_resources.assert_called_once_with()


def test_provider_cannot_declare_local_proof():
    from pydantic import ValidationError
    from qian_labor.ai.schemas import SourceLocator, ProviderSource
    with pytest.raises(ValidationError):
        SourceLocator(file_name="synthetic", _grounding={"status": "locally_located"})
    with pytest.raises(ValidationError):
        ProviderSource.model_validate({"file_name": "synthetic", "page": None, "row": None, "column": None,
            "sheet": None, "paragraph": None, "excerpt": "synthetic", "bbox": None,
            "provenance": "locally_located"})


def test_ten_employee_mixed_pipeline_sources_and_retry(tmp_path):
    import re
    from fastapi.testclient import TestClient
    from sqlalchemy import select
    from synthetic_mixed_materials import build_materials, EMPLOYEES, OCR_LINE
    from qian_labor.desktop.app import create_desktop_app
    from qian_labor.models.core import EmploymentFact, SourceLocator, UploadedFile, ProcessingJob
    from qian_labor.security.local_redaction import PrivacyBoundary, LocalImageRedactor, OCRToken
    from qian_labor.storage.local import LocalStorage
    from test_workspace_api import HEADERS, TOKEN
    class SyntheticOCR:
        calls = 0
        def extract_tokens(self, content):
            self.calls += 1
            return [OCRToken(OCR_LINE, 10, 20, 350, 12, "synthetic-line")]
    class EchoLocalText:
        # Deterministic provider fixture, not model acceptance or legal inference.
        name, is_external, calls = "offline-grounding-fixture", True, 0
        def extract(self, filename, content):
            self.calls += 1
            text = OCR_LINE if filename.endswith((".png", ".jpg")) else content.decode()
            facts = []
            for line in text.splitlines():
                if line.startswith(("[source", "[row_context")):
                    continue
                match = re.search(r"SYN-\d{3}", line)
                if match:
                    facts.append({"employee_id": match.group(), "fact_type": "employment.synthetic.observation", "value": line,
                        "confidence": 1, "source": {"file_name": "provider-invented-page-99", "excerpt": line}})
            return ExtractionResult.model_validate({"facts": facts})
    provider, ocr = EchoLocalText(), SyntheticOCR()
    app = create_desktop_app(data_dir=tmp_path / "app", launch_token=TOKEN)
    with TestClient(app) as client:
        aid = client.post("/api/analyses", headers=HEADERS, json={"name": "mixed synthetic", "company_display_name": "fictional"}).json()["id"]
        paths = []
        for name, content in build_materials().items():
            path = tmp_path / name
            path.write_bytes(content)
            paths.append(str(path))
        assert client.post(f"/api/analyses/{aid}/import-paths", headers=HEADERS, json={"paths": paths}).status_code == 200
        pipeline = ProcessingPipeline(app.state.database, LocalStorage(str(app.state.storage_root)), provider,
            privacy_boundary=PrivacyBoundary("synthetic-grounding-pepper-at-least-32", LocalImageRedactor(ocr)))
        pipeline.process(aid)
        assert ocr.calls == 3  # PNG + scan + the embedded DOCX image.
        with app.state.database.session() as session:
            sources = list(session.scalars(select(SourceLocator)))
            facts = list(session.scalars(select(EmploymentFact)))
            from qian_labor.services.source_provenance import deterministic_citation_id
            assert all(
                source.location.get("_citation_id") == deterministic_citation_id(
                    session.get(UploadedFile, source.file_id).sha256,
                    source.location,
                    source.excerpt,
                )
                for source in sources
            )
            assert set(EMPLOYEES) <= {re.search(r"SYN-\d{3}", f.value_json).group() for f in facts}
            assert all(s.location["_grounding"]["status"] == "locally_located" for s in sources)
            assert any(s.location.get("table") == 1 for s in sources)
            assert any(s.location.get("page") == 2 for s in sources)
            assert any(s.location.get("cell") == "A2" for s in sources)
            assert any(s.location.get("bbox") == [10, 20, 360, 32] for s in sources)
            embedded_file = next(f for f in session.scalars(select(UploadedFile))
                                 if f.original_filename == "embedded-warning.docx")
            embedded_sources = [s for s in sources if s.file_id == embedded_file.id]
            assert any(
                s.location.get("paragraph") == 2
                and s.location.get("image") == 1
                and "page" not in s.location
                for s in embedded_sources
            )
            before = {(s.id, s.fact_id, s.excerpt) for s in sources}
            file_ids = list(session.scalars(select(UploadedFile.id)))
            # Simulate an interrupted retry without deleting old facts/sources.
            for job in session.scalars(select(ProcessingJob).where(ProcessingJob.job_type == "extract")):
                job.status = "failed"
            session.commit()
        for fid in file_ids:
            pipeline._process_file(aid, fid)
        with app.state.database.session() as session:
            assert {(s.id, s.fact_id, s.excerpt) for s in session.scalars(select(SourceLocator))} == before


def test_oversize_file_makes_no_call_and_other_success_is_kept(tmp_path):
    from fastapi.testclient import TestClient
    from qian_labor.desktop.app import create_desktop_app
    from qian_labor.storage.local import LocalStorage
    from test_workspace_api import HEADERS, TOKEN
    class Provider:
        name, is_external, calls = "fake", False, 0
        def extract(self, filename, content):
            self.calls += 1
            return result("SYN-001 signed contract")
    provider = Provider()
    app = create_desktop_app(data_dir=tmp_path / "app", launch_token=TOKEN)
    with TestClient(app) as client:
        aid = client.post("/api/analyses", headers=HEADERS, json={"name": "synthetic", "company_display_name": "fictional"}).json()["id"]
        good, large = tmp_path / "good.docx", tmp_path / "large.docx"
        good.write_bytes(word_bytes())
        doc = Document()
        doc.add_paragraph("x" * 100_001)
        doc.save(large)
        client.post(f"/api/analyses/{aid}/import-paths", headers=HEADERS, json={"paths": [str(good), str(large)]})
        outcome = ProcessingPipeline(app.state.database, LocalStorage(str(app.state.storage_root)), provider).process(aid)
        assert provider.calls == 1 and outcome["status"] == "partial"
        workspace = client.get(f"/api/analyses/{aid}/workspace", headers=HEADERS).json()
        assert workspace["files"][0]["fact_count"] == 1
        assert workspace["files"][1]["error_code"] == "AI_TEXT_LIMIT_EXCEEDED"


def test_empty_docx_does_not_fallback_to_sending_raw_office_bytes():
    from qian_labor.ai.providers import AIProviderError
    doc = Document()
    output = BytesIO()
    doc.save(output)
    parsed = ParserRegistry().parse("empty.docx", output.getvalue())
    with pytest.raises(AIProviderError, match="AI_NO_SUPPORTED_FACTS"):
        ProcessingPipeline._extraction_inputs("empty.docx", output.getvalue(), parsed)


def test_empty_extraction_is_not_upgraded_success_using_old_facts(tmp_path):
    from fastapi.testclient import TestClient
    from sqlalchemy import select
    from qian_labor.desktop.app import create_desktop_app
    from qian_labor.models.core import EmploymentFact, UploadedFile
    from qian_labor.storage.local import LocalStorage
    from test_workspace_api import HEADERS, TOKEN
    class EmptyProvider:
        name, is_external = "fake", False
        def extract(self, filename, content):
            return ExtractionResult()
    app = create_desktop_app(data_dir=tmp_path / "app", launch_token=TOKEN)
    with TestClient(app) as client:
        aid = client.post("/api/analyses", headers=HEADERS, json={"name": "synthetic", "company_display_name": "fictional"}).json()["id"]
        path = tmp_path / "old.docx"
        path.write_bytes(word_bytes())
        client.post(f"/api/analyses/{aid}/import-paths", headers=HEADERS, json={"paths": [str(path)]})
        with app.state.database.session() as session:
            file = session.scalar(select(UploadedFile))
            session.add(EmploymentFact(analysis_id=aid, file_id=file.id, fact_type="employment.contract.exists", value_json=True,
                normalized_value_json=True, verification_status="confirmed", extraction_method="old", confidence=1, dedupe_key="old"))
            session.commit()
        pipeline = ProcessingPipeline(app.state.database, LocalStorage(str(app.state.storage_root)), EmptyProvider())
        outcome = pipeline.process(aid)
        assert outcome["files"][0]["error_code"] == "AI_NO_SUPPORTED_FACTS"
        workspace = client.get(f"/api/analyses/{aid}/workspace", headers=HEADERS).json()
        assert workspace["files"][0]["extraction_version"] is None
        assert workspace["files"][0]["fact_count"] == 1


def test_cross_file_same_value_keeps_owned_facts_and_unlocated_uncertainty(review_case, tmp_path):
    from sqlalchemy import select
    from qian_labor.models.core import EmploymentFact, SourceLocator, UploadedFile
    from qian_labor.services.risk_evaluation import RiskEvaluationService
    from qian_labor.services.source_provenance import uncertain_grounded_fact_ids
    from qian_labor.storage.local import LocalStorage
    _, database, aid, _, _ = review_case
    pipeline = ProcessingPipeline(database, LocalStorage(tmp_path))
    with database.session() as session:
        first = session.scalar(select(UploadedFile))
        second = UploadedFile(analysis_id=aid, original_filename="synthetic-second.docx", storage_key="synthetic/second.docx",
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", extension=".docx", size_bytes=1, sha256="b" * 64)
        session.add(second)
        session.flush()
        for file, state in [(first, "locally_located"), (second, "unlocated_needs_review")]:
            extracted = result("SYN-001 signed contract")
            if state == "unlocated_needs_review":
                extracted.facts[0].source.excerpt = ""
            pipeline._persist_result(session, aid, file, extracted, grounding=[{
                "version": "parser-grounding-v3", "status": state, "requires_review": False}])
        session.flush()
        facts = list(session.scalars(select(EmploymentFact).where(EmploymentFact.analysis_id == aid,
            EmploymentFact.extraction_method == pipeline.provider.name)))
        assert len(facts) == 2 and {f.file_id for f in facts} == {first.id, second.id}
        assert facts[0].value_json is True and facts[1].value_json is True
        sources = list(session.scalars(select(SourceLocator).where(SourceLocator.fact_id.in_([f.id for f in facts]))))
        assert all(s.file_id == next(f.file_id for f in facts if f.id == s.fact_id) for s in sources)
        assert uncertain_grounded_fact_ids(session, aid) == {f.id for f in facts if f.file_id == second.id}
        assert RiskEvaluationService._context_facts(session, facts)["employment.contract.exists"].conflicted
