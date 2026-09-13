from pathlib import Path

from fastapi.testclient import TestClient
from sqlalchemy import select

from qian_labor.ai.schemas import EmploymentFact, ExtractionResult, SourceLocator
from qian_labor.ai.grounding import ground_result
from qian_labor.desktop.app import create_desktop_app
from qian_labor.jobs.processing import ProcessingPipeline
from qian_labor.models.core import SourceLocator as StoredSource, UploadedFile
from qian_labor.parsers.protocols import ParsedBlock, ParsedDocument
from qian_labor.parsers.registry import ParserRegistry
from qian_labor.storage.local import LocalStorage


def result(employee_id: str | None = "QA-907") -> ExtractionResult:
    return ExtractionResult(facts=[EmploymentFact(
        employee_id=employee_id,
        fact_type="employment.contract.exists",
        value=False,
        confidence=0.9,
        source=SourceLocator(file_name="invented.csv", row=1, column="wrong", excerpt="模型改写的说明"),
    )])


def test_provider_input_keeps_parser_owned_spreadsheet_coordinates() -> None:
    content = "员工编号,合同存在\nQA-907,false\n".encode()
    parsed = ParserRegistry().parse("synthetic.csv", content)
    text = ProcessingPipeline._extraction_inputs("synthetic.csv", content, parsed)[0][1].decode()
    assert '"sheet": "CSV"' in text
    assert '"row": 2' in text


def test_model_row_one_is_rebound_to_actual_data_row_and_original_excerpt() -> None:
    parsed = ParserRegistry().parse("synthetic.csv", "员工编号,合同存在\nQA-907,false\n".encode())
    extracted = result()
    ProcessingPipeline._bind_spreadsheet_sources(extracted, parsed, "synthetic.csv")
    source = extracted.facts[0].source
    assert (source.file_name, source.sheet, source.row) == ("synthetic.csv", "CSV", 2)
    assert source.excerpt == "QA-907 | false"
    assert source.column is None


def test_employee_identifier_selects_actual_row_not_model_claim_or_substring() -> None:
    parsed = ParserRegistry().parse("synthetic.csv", "员工编号,合同存在\nQA-9070,true\nQA-907,false\n".encode())
    extracted = result()
    ProcessingPipeline._bind_spreadsheet_sources(extracted, parsed, "synthetic.csv")
    assert extracted.facts[0].source.row == 3


def test_ambiguous_rows_do_not_publish_unverified_coordinates() -> None:
    parsed = ParserRegistry().parse("synthetic.csv", "员工编号,合同存在\nQA-907,true\nQA-907,false\n".encode())
    extracted = result()
    ProcessingPipeline._bind_spreadsheet_sources(extracted, parsed, "synthetic.csv")
    assert extracted.facts[0].source.row is None
    assert extracted.facts[0].source.excerpt == ""
    assert extracted.facts[0].needs_human_confirmation


def test_excel_cells_are_grouped_by_sheet_and_row() -> None:
    parsed = ParsedDocument("spreadsheet", [
        ParsedBlock("QA-9070", "cell", {"sheet": "工资", "row": 2, "column": 1}),
        ParsedBlock("QA-907", "cell", {"sheet": "合同", "row": 7, "column": 1}),
        ParsedBlock("false", "cell", {"sheet": "合同", "row": 7, "column": 2}),
    ])
    extracted = result()
    ProcessingPipeline._bind_spreadsheet_sources(extracted, parsed, "synthetic.xlsx")
    assert (extracted.facts[0].source.sheet, extracted.facts[0].source.row) == ("合同", 7)
    assert extracted.facts[0].source.excerpt == "QA-907 | false"


def test_repeated_value_with_full_row_excerpt_binds_to_employee_row() -> None:
    """A model may quote a whole spreadsheet row rather than one cell.

    The date is intentionally repeated for two employees.  The local row
    identity, not the repeated value alone, must select the source row.
    """
    parsed = ParsedDocument("spreadsheet", [
        ParsedBlock("QA-401", "cell", {"sheet": "trial", "row": 2, "column": 1}),
        ParsedBlock("2026-03-01", "cell", {"sheet": "trial", "row": 2, "column": 3}),
        ParsedBlock("QA-402", "cell", {"sheet": "trial", "row": 3, "column": 1}),
        ParsedBlock("2026-03-01", "cell", {"sheet": "trial", "row": 3, "column": 3}),
    ])
    extracted = ExtractionResult.model_validate({
        "facts": [{
            "employee_id": "QA-402",
            "fact_type": "employment.probation.end_date",
            "value": "2026-03-01",
            "confidence": 0.9,
            "source": {
                "file_name": "model.xlsx",
                "excerpt": "QA-402 | 2026-03-01",
            },
        }],
    })
    item = ProcessingPipeline._extraction_inputs("synthetic.xlsx", b"", parsed)[0]

    proofs = ground_result(extracted, item, "synthetic.xlsx")

    assert proofs[0]["status"] == "locally_located"
    assert extracted.facts[0].source.sheet == "trial"
    assert extracted.facts[0].source.row == 3


def test_repeated_dates_remain_scoped_to_exact_employee_row() -> None:
    """A repeated value must not bind to a different employee's row."""
    from qian_labor.ai.grounding import ExtractionInput, ground_result

    blocks = (
        ParsedBlock("QA-401", "cell", {"sheet": "trial", "row": 2, "column": 1}),
        ParsedBlock("2026-03-01", "cell", {"sheet": "trial", "row": 2, "column": 3}),
        ParsedBlock("QA-402", "cell", {"sheet": "trial", "row": 3, "column": 1}),
        ParsedBlock("2026-03-01", "cell", {"sheet": "trial", "row": 3, "column": 3}),
    )
    extracted = ExtractionResult(facts=[EmploymentFact(
        employee_id="QA-402",
        fact_type="employment.probation.end_date",
        value="2026-03-01",
        confidence=0.9,
        source=SourceLocator(file_name="synthetic.xlsx", excerpt="2026-03-01"),
    )])

    proofs = ground_result(
        extracted,
        ExtractionInput("synthetic.xlsx", b"", blocks),
        "synthetic.xlsx",
    )

    assert len(extracted.facts) == 1
    assert extracted.facts[0].source.row == 3
    assert extracted.facts[0].source.column == "3"
    assert proofs[0]["status"] == "locally_located"


def test_equal_row_numbers_on_different_sheets_are_not_merged() -> None:
    """Sheet identity is part of a spreadsheet row's parser-owned scope."""
    from qian_labor.ai.grounding import ExtractionInput, ground_result

    blocks = (
        ParsedBlock("QA-402", "cell", {"sheet": "trial", "row": 3, "column": 1}),
        ParsedBlock("2026-03-01", "cell", {"sheet": "trial", "row": 3, "column": 3}),
        ParsedBlock("QA-402", "cell", {"sheet": "archive", "row": 3, "column": 1}),
        ParsedBlock("2026-03-01", "cell", {"sheet": "archive", "row": 3, "column": 3}),
    )
    extracted = ExtractionResult(facts=[EmploymentFact(
        employee_id="QA-402",
        fact_type="employment.probation.end_date",
        value="2026-03-01",
        confidence=0.9,
        source=SourceLocator(file_name="synthetic.xlsx", sheet="trial", excerpt="2026-03-01"),
    )])

    proofs = ground_result(
        extracted,
        ExtractionInput("synthetic.xlsx", b"", blocks),
        "synthetic.xlsx",
    )

    assert extracted.facts[0].source.sheet == "trial"
    assert extracted.facts[0].source.row == 3
    assert extracted.facts[0].source.column == "3"
    assert proofs[0]["status"] == "locally_located"


def test_employee_prefix_collision_does_not_cross_bind_rows() -> None:
    """QA-907 must not match the longer QA-9070 identifier."""
    from qian_labor.ai.grounding import ExtractionInput, ground_result

    blocks = (
        ParsedBlock("QA-9070", "cell", {"sheet": "trial", "row": 2, "column": 1}),
        ParsedBlock("2026-02-01", "cell", {"sheet": "trial", "row": 2, "column": 3}),
        ParsedBlock("QA-907", "cell", {"sheet": "trial", "row": 3, "column": 1}),
        ParsedBlock("2026-03-01", "cell", {"sheet": "trial", "row": 3, "column": 3}),
    )
    extracted = ExtractionResult(facts=[EmploymentFact(
        employee_id="QA-907",
        fact_type="employment.probation.end_date",
        value="2026-03-01",
        confidence=0.9,
        source=SourceLocator(file_name="synthetic.xlsx", excerpt="2026-03-01"),
    )])

    proofs = ground_result(
        extracted,
        ExtractionInput("synthetic.xlsx", b"", blocks),
        "synthetic.xlsx",
    )

    assert extracted.facts[0].source.row == 3
    assert extracted.facts[0].source.column == "3"
    assert proofs[0]["status"] == "locally_located"


def test_explicit_wrong_spreadsheet_coordinates_are_rejected() -> None:
    """A false model row/column hint cannot be silently repaired."""
    from qian_labor.ai.grounding import ExtractionInput, ground_result

    blocks = (
        ParsedBlock("QA-402", "cell", {"sheet": "trial", "row": 3, "column": 1}),
        ParsedBlock("2026-03-01", "cell", {"sheet": "trial", "row": 3, "column": 3}),
    )
    extracted = ExtractionResult(facts=[EmploymentFact(
        employee_id="QA-402",
        fact_type="employment.probation.end_date",
        value="2026-03-01",
        confidence=0.9,
        source=SourceLocator(
            file_name="synthetic.xlsx",
            sheet="trial",
            row=99,
            column="9",
            excerpt="2026-03-01",
        ),
    )])

    proofs = ground_result(
        extracted,
        ExtractionInput("synthetic.xlsx", b"", blocks),
        "synthetic.xlsx",
    )

    assert proofs[0]["status"] == "unlocated_needs_review"
    assert extracted.facts[0].source.row is None
    assert extracted.facts[0].source.excerpt == ""
    assert extracted.facts[0].needs_human_confirmation


def test_duplicate_parser_locations_remain_pending_review() -> None:
    """Duplicate parser evidence is explicit ambiguity, not a single proof."""
    from qian_labor.ai.grounding import ExtractionInput, ground_result

    blocks = (
        ParsedBlock("QA-402", "cell", {"sheet": "trial", "row": 3, "column": 1}),
        ParsedBlock("2026-03-01", "cell", {"sheet": "trial", "row": 3, "column": 3}),
        ParsedBlock("2026-03-01", "cell", {"sheet": "trial", "row": 3, "column": 4}),
    )
    extracted = ExtractionResult(facts=[EmploymentFact(
        employee_id="QA-402",
        fact_type="employment.probation.end_date",
        value="2026-03-01",
        confidence=0.9,
        source=SourceLocator(file_name="synthetic.xlsx", excerpt="2026-03-01"),
    )])

    proofs = ground_result(
        extracted,
        ExtractionInput("synthetic.xlsx", b"", blocks),
        "synthetic.xlsx",
    )

    assert len(extracted.facts) == 2
    assert all(f.source.row == 3 for f in extracted.facts)
    assert all(f.needs_human_confirmation for f in extracted.facts)
    assert all(p["status"] == "locally_located" for p in proofs)


def test_pipeline_locates_actual_quote_without_trusting_provider_filename(tmp_path: Path) -> None:
    class ActualQuoteProvider:
        name = "source-binding-test"
        is_external = False

        def extract(self, filename: str, content: bytes) -> ExtractionResult:
            extracted = result()
            extracted.facts[0].source = SourceLocator(file_name="invented.csv", excerpt="QA-907 | false")
            return extracted

    token = "synthetic-source-binding-token"
    headers = {"X-Qian-Desktop-Token": token}
    source = tmp_path / "synthetic.csv"
    source.write_text("员工编号,合同存在\nQA-907,false\n", encoding="utf-8")
    app = create_desktop_app(data_dir=tmp_path / "data", launch_token=token)
    with TestClient(app) as client:
        analysis_id = client.post("/api/analyses", headers=headers, json={
            "name": "虚构来源核验", "company_display_name": "虚构企业",
        }).json()["id"]
        assert client.post(f"/api/analyses/{analysis_id}/import-paths", headers=headers,
                           json={"paths": [str(source)]}).status_code == 200
        ProcessingPipeline(app.state.database, LocalStorage(str(app.state.storage_root)),
                           ActualQuoteProvider()).process(analysis_id)
        with app.state.database.session() as session:
            stored = session.scalar(select(StoredSource).where(StoredSource.analysis_id == analysis_id))
            assert stored is not None
            from qian_labor.services.source_provenance import deterministic_citation_id
            from qian_labor.ai.grounding import EXTRACTION_VERSION
            assert stored.location == {"sheet": "CSV", "row": 2, "_grounding": {
                "version": EXTRACTION_VERSION, "status": "locally_located", "requires_review": True},
                "_citation_id": deterministic_citation_id(
                    session.scalar(select(UploadedFile.sha256).where(UploadedFile.analysis_id == analysis_id)),
                    {"sheet": "CSV", "row": 2, "_grounding": {
                        "version": EXTRACTION_VERSION, "status": "locally_located", "requires_review": True}},
                    "QA-907 | false")}
            assert stored.excerpt == "QA-907 | false"


def test_pipeline_persists_located_cell_with_recomputable_citation(tmp_path: Path) -> None:
    import io

    from openpyxl import Workbook

    class CellQuoteProvider:
        name = "cell-source-binding-test"
        is_external = False

        def extract(self, filename: str, content: bytes) -> ExtractionResult:
            extracted = result()
            extracted.facts[0].source = SourceLocator(
                file_name="invented.xlsx", sheet="试用期", row=2, column="2", excerpt="False"
            )
            return extracted

    token = "synthetic-cell-source-binding-token"
    headers = {"X-Qian-Desktop-Token": token}
    source = tmp_path / "synthetic.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "试用期"
    sheet.append(["员工编号", "合同存在"])
    sheet.append(["QA-907", False])
    buffer = io.BytesIO()
    book.save(buffer)
    source.write_bytes(buffer.getvalue())
    app = create_desktop_app(data_dir=tmp_path / "data", launch_token=token)
    with TestClient(app) as client:
        analysis_id = client.post("/api/analyses", headers=headers, json={
            "name": "虚构单元格来源核验", "company_display_name": "虚构企业",
        }).json()["id"]
        assert client.post(f"/api/analyses/{analysis_id}/import-paths", headers=headers,
                           json={"paths": [str(source)]}).status_code == 200
        ProcessingPipeline(app.state.database, LocalStorage(str(app.state.storage_root)),
                           CellQuoteProvider()).process(analysis_id)
        with app.state.database.session() as session:
            stored = session.scalar(select(StoredSource).where(StoredSource.analysis_id == analysis_id))
            uploaded = session.scalar(select(UploadedFile).where(UploadedFile.analysis_id == analysis_id))
            assert stored is not None and uploaded is not None
            assert stored.excerpt == "False"
            assert stored.location["sheet"] == "试用期"
            assert stored.location["row"] == 2
            assert stored.location["column"] == "2"
            assert stored.location["_grounding"]["status"] == "locally_located"
            citation_location = {key: value for key, value in stored.location.items()
                                 if key != "_citation_id"}
            from qian_labor.services.source_provenance import deterministic_citation_id
            assert stored.location["_citation_id"] == deterministic_citation_id(
                uploaded.sha256, citation_location, stored.excerpt
            )
