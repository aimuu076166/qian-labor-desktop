"""Ten-person mixed-file acceptance fixture, with a deterministic provider.

This checks import/parser/matching/persistence integration, NOT real GLM quality.
All generated files live in pytest's temporary directory; no customer data.
"""

from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image, ImageDraw
from sqlalchemy import select

from qian_labor.ai.schemas import EmploymentFact, ExtractionResult, SourceLocator
from qian_labor.desktop.app import create_desktop_app
from qian_labor.jobs.processing import ProcessingPipeline
from qian_labor.models.core import Employee, EmploymentFact as StoredFact
from qian_labor.storage.local import LocalStorage

NUMBERS = tuple(f"SYN-{index:03d}" for index in range(1, 11))
TOKEN = "synthetic-ten-person-workflow-token"
HEADERS = {"X-Qian-Desktop-Token": TOKEN}


def mixed_materials(root: Path) -> list[Path]:
    roster = root / "synthetic-roster.csv"
    roster.write_text("工号,姓名,在职状态\n" + "".join(
        f"{number},完全虚构员工{index},在职\n" for index, number in enumerate(NUMBERS, 1)
    ), encoding="utf-8")
    paths = [roster]
    for month, wage in (("2026-01", 5000), ("2026-02", 4500)):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "工资"
        sheet.append(["工号", "工资月份", "实发工资"])
        sheet.append(["SYN-010", month, wage])
        path = root / f"synthetic-payroll-{month}.xlsx"
        workbook.save(path)
        paths.append(path)
    document = Document()
    document.add_paragraph("完全虚构合同。工号 SYN-010，约定月工资 5000 元。仅供测试。")
    contract = root / "synthetic-contract.docx"
    document.save(contract)
    paths.append(contract)
    scanned = root / "synthetic-contract-scan.png"
    with Image.new("RGB", (900, 160), "white") as image:
        ImageDraw.Draw(image).text((20, 40), "SYNTHETIC ONLY: SYN-010 contract wage 5000", fill="black")
        image.save(scanned)
    paths.append(scanned)
    return paths


class ControlledExtraction:
    name = "ten-person-controlled-extraction"
    is_external = False

    def extract(self, filename: str, content: bytes) -> ExtractionResult:
        if filename.endswith("roster.csv"):
            return ExtractionResult(document_type="roster", facts=[EmploymentFact(
                employee_id=number, fact_type="employment.status", value="active",
                confidence=1, source=SourceLocator(file_name=filename, row=index + 1,
                                                 excerpt=f"{number} 在职"),
            ) for index, number in enumerate(NUMBERS, 1)])
        if "payroll" in filename:
            month = "2026-01" if "2026-01" in filename else "2026-02"
            return ExtractionResult(document_type="payroll", employee_number="SYN-010", facts=[
                EmploymentFact(employee_id="SYN-010", fact_type="employment.pay.actual_wage",
                               value=5000 if month == "2026-01" else 4500, confidence=1,
                               source=SourceLocator(file_name=filename, row=2, sheet="工资",
                                                    excerpt=f"SYN-010 {month}")),
            ])
        return ExtractionResult(document_type="contract", employee_number="SYN-010", facts=[
            EmploymentFact(employee_id="SYN-010", fact_type="employment.pay.contract_wage",
                           value=5000, confidence=1, source=SourceLocator(file_name=filename,
                           page=1 if filename.endswith("png") else None,
                           paragraph=None if filename.endswith("png") else 1,
                           excerpt="SYN-010 contract wage 5000")),
        ])


def test_ten_people_keep_confirmed_identity_when_mixed_materials_are_added(tmp_path: Path) -> None:
    paths = mixed_materials(tmp_path)
    app = create_desktop_app(data_dir=tmp_path / "data", launch_token=TOKEN)
    with TestClient(app) as client:
        analysis_id = client.post("/api/analyses", headers=HEADERS, json={
            "name": "十人混合材料回归", "company_display_name": "完全虚构测试企业",
        }).json()["id"]
        def import_paths(selected: list[Path]) -> None:
            response = client.post(f"/api/analyses/{analysis_id}/import-paths", headers=HEADERS,
                                   json={"paths": [str(path) for path in selected]})
            assert response.status_code == 200
        def process() -> None:
            ProcessingPipeline(app.state.database, LocalStorage(str(app.state.storage_root)),
                               ControlledExtraction()).process(analysis_id)
        import_paths(paths[:1])
        process()
        candidates = client.get(f"/api/analyses/{analysis_id}/matching-candidates",
                                headers=HEADERS).json()["candidates"]
        assert len(candidates) == 10
        for candidate in candidates:
            number = candidate["extracted_fields"]["employee_ids"][0]
            response = client.post(f"/api/analyses/{analysis_id}/matching-decisions",
                headers=HEADERS, json={"candidate_id": candidate["id"],
                    "decision": "create_unknown", "display_name": f"虚构员工{number}",
                    "employee_number": number, "fact_ids": candidate["fact_ids"]})
            assert response.status_code == 200
        import_paths(paths[1:])
        process()
        import_paths(paths[1:])  # Duplicate import must retain one material copy.
        with app.state.database.session() as session:
            employees = list(session.scalars(select(Employee).where(Employee.analysis_id == analysis_id)))
            assert len(employees) == 10
            assert {employee.employee_number for employee in employees} == set(NUMBERS)
            target_id = next(employee.id for employee in employees if employee.employee_number == "SYN-010")
            wage_facts = list(session.scalars(select(StoredFact).where(
                StoredFact.analysis_id == analysis_id,
                StoredFact.fact_type.in_(["employment.pay.actual_wage", "employment.pay.contract_wage"]),
            )))
            assert len(wage_facts) == 4
            assert {fact.employee_id for fact in wage_facts} == {target_id}
        status = client.get(f"/api/analyses/{analysis_id}/processing", headers=HEADERS).json()
        assert len(status["files"]) == 5
        assert all(item["status"] == "processed" for item in status["files"])
        ledger = client.get(f"/api/analyses/{analysis_id}/employees", headers=HEADERS).json()
        assert ledger["total"] == 10
        assert len(ledger["items"]) == 10
